#!/usr/bin/env python3
"""Build the official document DOCX template library."""

from __future__ import annotations

import json
import re
import shutil
import tempfile
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_DIR = ROOT / "assets" / "templates"

FONT_TITLE = "Hiragino Mincho ProN"
FONT_BODY = "Hiragino Mincho ProN"
FONT_HEITI = "Hiragino Sans GB"
FONT_KAITI = "Hiragino Mincho ProN"
RED = "C00000"
STYLE_TITLE = "公文标题"
STYLE_BODY = "正文"
STYLE_H1 = "一级标题"
STYLE_H2 = "二级标题"
STYLE_H3 = "三级标题"
STYLE_META = "元信息"
STYLE_SIGNATURE = "落款"
STYLE_COPY_TO = "版记"


def set_run_font(run, font_name: str, size_pt: float | None = None, bold: bool | None = None, color: str = "000000") -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_font(paragraph, font_name: str, size_pt: float, bold: bool | None = None, color: str = "000000") -> None:
    for run in paragraph.runs:
        set_run_font(run, font_name, size_pt, bold, color)


def set_paragraph_rhythm(paragraph, first_line_chars: float = 0, line_pt: float = 29, before_pt: float = 0, after_pt: float = 0) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.line_spacing = Pt(line_pt)
    fmt.space_before = Pt(before_pt)
    fmt.space_after = Pt(after_pt)
    if first_line_chars:
        fmt.first_line_indent = Pt(16 * first_line_chars)


def paragraph_border(paragraph, bottom: str | None = None, top: str | None = None, size: int = 12, space: int = 1) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    for edge, color in (("top", top), ("bottom", bottom)):
        if not color:
            continue
        elem = OxmlElement(f"w:{edge}")
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), str(size))
        elem.set(qn("w:space"), str(space))
        elem.set(qn("w:color"), color)
        p_bdr.append(elem)


def set_cell_text(cell, text: str, font_name: str = FONT_BODY, size_pt: float = 14, align=WD_ALIGN_PARAGRAPH.LEFT) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.text = text
    set_paragraph_rhythm(p, line_pt=24)
    set_paragraph_font(p, font_name, size_pt)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tbl_cell_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_style(document: Document, style_id: str, name: str, font: str, size: float, bold: bool = False, color: str = "000000") -> None:
    styles = document.styles
    style = styles.add_style(style_id, 1)
    style.name = name
    style.font.name = font
    style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    style.paragraph_format.line_spacing = Pt(29)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3.7)
    section.bottom_margin = Cm(3.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(2.5)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT_BODY
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)
    normal.font.size = Pt(16)
    normal.font.color.rgb = RGBColor.from_string("000000")
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(29)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    add_style(document, "DocTitle", "公文标题", FONT_TITLE, 22, False)
    add_style(document, "Body", "正文", FONT_BODY, 16, False)
    add_style(document, "Heading1", "一级标题", FONT_HEITI, 16, False)
    add_style(document, "Heading2", "二级标题", FONT_KAITI, 16, False)
    add_style(document, "Heading3", "三级标题", FONT_BODY, 16, True)
    add_style(document, "Meta", "元信息", FONT_BODY, 14, False)
    add_style(document, "Signature", "落款", FONT_BODY, 16, False)
    add_style(document, "CopyTo", "版记", FONT_BODY, 12, False)


def add_p(document: Document, text: str = "", style: str = STYLE_BODY, align=None, first_line: float = 2, line_pt: float = 29):
    p = document.add_paragraph(style=style)
    if text:
        p.add_run(text)
    if align is not None:
        p.alignment = align
    set_paragraph_rhythm(p, first_line_chars=first_line, line_pt=line_pt)
    return p


def add_spacer(document: Document, line_pt: float = 12) -> None:
    p = document.add_paragraph()
    set_paragraph_rhythm(p, line_pt=line_pt)


def add_page_number(section) -> None:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("— ")
    set_run_font(run, FONT_BODY, 14)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    run2 = p.add_run(" —")
    set_run_font(run2, FONT_BODY, 14)


def patch_ooxml(path: Path) -> None:
    with ZipFile(path, "r") as src:
        replacements: dict[str, str] = {}
        if "word/settings.xml" in src.namelist():
            settings = src.read("word/settings.xml").decode("utf-8")
            replacements["word/settings.xml"] = settings.replace(
                'w:name="compatibilityMode" w:uri="http://schemas.microsoft.com/office/word" w:val="14"',
                'w:name="compatibilityMode" w:uri="http://schemas.microsoft.com/office/word" w:val="15"',
            )
        if "word/styles.xml" in src.namelist():
            styles = src.read("word/styles.xml").decode("utf-8")
            styles = re.sub(
                r'<w:color w:val="(?:4F81BD|5B9BD5|2F5496|1F4E79)"(?: w:themeColor="[^"]+")?/>',
                '<w:color w:val="000000"/>',
                styles,
            )
            styles = re.sub(
                r'w:color="(?:4F81BD|5B9BD5|2F5496|1F4E79)"(?: w:themeColor="[^"]+")?',
                'w:color="000000"',
                styles,
            )
            styles = re.sub(
                r'w:fill="(?:4F81BD|5B9BD5|2F5496|1F4E79)"(?: w:themeFill="[^"]+")?',
                'w:fill="E6E6E6"',
                styles,
            )
            replacements["word/styles.xml"] = styles
        if not replacements:
            return
        with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
            tmp_path = Path(tmp.name)
        try:
            with ZipFile(tmp_path, "w") as dst:
                for item in src.infolist():
                    if item.filename in replacements:
                        dst.writestr(item, replacements[item.filename])
                    else:
                        dst.writestr(item, src.read(item.filename))
            shutil.move(str(tmp_path), path)
        finally:
            if tmp_path.exists():
                tmp_path.unlink()
    path.chmod(0o644)


def official_document(path: Path) -> None:
    doc = Document()
    configure_document(doc)
    add_page_number(doc.sections[0])

    for text in ["{{份号}}", "{{密级和保密期限}}", "{{紧急程度}}"]:
        p = add_p(doc, text, STYLE_META, first_line=0, line_pt=24)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    add_spacer(doc, 16)
    p = add_p(doc, "{{发文机关标志}}文件", STYLE_TITLE, WD_ALIGN_PARAGRAPH.CENTER, first_line=0, line_pt=34)
    set_paragraph_font(p, FONT_TITLE, 28, False, RED)
    add_spacer(doc, 14)

    p = add_p(doc, "{{发文字号}}", STYLE_META, WD_ALIGN_PARAGRAPH.CENTER, first_line=0, line_pt=24)
    paragraph_border(p, bottom=RED, size=12, space=8)

    add_spacer(doc, 18)
    p = add_p(doc, "{{标题}}", STYLE_TITLE, WD_ALIGN_PARAGRAPH.CENTER, first_line=0, line_pt=34)
    set_paragraph_font(p, FONT_TITLE, 22)
    add_spacer(doc, 8)

    add_p(doc, "{{主送机关}}：", STYLE_BODY, first_line=0)
    add_p(doc, "{{正文第一段}}", STYLE_BODY)
    add_p(doc, "一、{{一级标题}}", STYLE_H1)
    add_p(doc, "{{正文第二段}}", STYLE_BODY)
    add_p(doc, "（一）{{二级标题}}", STYLE_H2)
    add_p(doc, "{{正文第三段}}", STYLE_BODY)
    add_p(doc, "附件：{{附件名称}}", STYLE_BODY, first_line=0)
    add_spacer(doc, 24)

    p = add_p(doc, "{{发文机关署名}}", STYLE_SIGNATURE, WD_ALIGN_PARAGRAPH.RIGHT, first_line=0)
    p.paragraph_format.right_indent = Cm(1.0)
    p = add_p(doc, "{{成文日期}}", STYLE_SIGNATURE, WD_ALIGN_PARAGRAPH.RIGHT, first_line=0)
    p.paragraph_format.right_indent = Cm(1.0)

    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_cell_margins(table)
    set_cell_text(table.cell(0, 0), "抄送：{{抄送机关}}。", FONT_BODY, 12)
    set_cell_text(table.cell(0, 1), "", FONT_BODY, 12)
    set_cell_text(table.cell(1, 0), "{{印发机关}}", FONT_BODY, 12)
    set_cell_text(table.cell(1, 1), "{{印发日期}}印发", FONT_BODY, 12, WD_ALIGN_PARAGRAPH.RIGHT)
    doc.save(path)
    patch_ooxml(path)


def official_letter(path: Path) -> None:
    doc = Document()
    configure_document(doc)
    add_page_number(doc.sections[0])

    p = add_p(doc, "{{发文机关标志}}", STYLE_TITLE, WD_ALIGN_PARAGRAPH.CENTER, first_line=0, line_pt=34)
    set_paragraph_font(p, FONT_TITLE, 26, False, RED)
    p = add_p(doc, "", STYLE_META, WD_ALIGN_PARAGRAPH.CENTER, first_line=0, line_pt=8)
    paragraph_border(p, top=RED, bottom=RED, size=10, space=2)
    add_spacer(doc, 22)

    p = add_p(doc, "{{标题}}", STYLE_TITLE, WD_ALIGN_PARAGRAPH.CENTER, first_line=0, line_pt=34)
    set_paragraph_font(p, FONT_TITLE, 22)
    add_spacer(doc, 8)
    add_p(doc, "{{主送机关}}：", STYLE_BODY, first_line=0)
    add_p(doc, "{{正文第一段}}", STYLE_BODY)
    add_p(doc, "一、{{一级标题}}", STYLE_H1)
    add_p(doc, "{{正文第二段}}", STYLE_BODY)
    add_p(doc, "二、{{二级标题}}", STYLE_H1)
    add_p(doc, "{{正文第三段}}", STYLE_BODY)
    add_p(doc, "联系人：{{联系人}}，联系电话：{{联系电话}}。", STYLE_BODY)
    add_spacer(doc, 24)
    p = add_p(doc, "{{发文机关署名}}", STYLE_SIGNATURE, WD_ALIGN_PARAGRAPH.RIGHT, first_line=0)
    p.paragraph_format.right_indent = Cm(1.0)
    p = add_p(doc, "{{成文日期}}", STYLE_SIGNATURE, WD_ALIGN_PARAGRAPH.RIGHT, first_line=0)
    p.paragraph_format.right_indent = Cm(1.0)
    doc.save(path)
    patch_ooxml(path)


def meeting_material(path: Path) -> None:
    doc = Document()
    configure_document(doc)
    add_page_number(doc.sections[0])

    p = add_p(doc, "{{会议名称}}", STYLE_META, WD_ALIGN_PARAGRAPH.CENTER, first_line=0, line_pt=24)
    set_paragraph_font(p, FONT_KAITI, 16)
    p = add_p(doc, "{{标题}}", STYLE_TITLE, WD_ALIGN_PARAGRAPH.CENTER, first_line=0, line_pt=34)
    set_paragraph_font(p, FONT_TITLE, 22)
    p = add_p(doc, "{{讲话人}}  {{日期}}", STYLE_META, WD_ALIGN_PARAGRAPH.CENTER, first_line=0, line_pt=24)
    set_paragraph_font(p, FONT_BODY, 14)
    add_spacer(doc, 16)
    add_p(doc, "同志们：", STYLE_BODY, first_line=0)
    add_p(doc, "{{开场段}}", STYLE_BODY)
    add_p(doc, "一、{{第一部分标题}}", STYLE_H1)
    add_p(doc, "{{第一部分正文}}", STYLE_BODY)
    add_p(doc, "二、{{第二部分标题}}", STYLE_H1)
    add_p(doc, "{{第二部分正文}}", STYLE_BODY)
    add_p(doc, "三、{{第三部分标题}}", STYLE_H1)
    add_p(doc, "{{第三部分正文}}", STYLE_BODY)
    add_p(doc, "{{结束段}}", STYLE_BODY)
    doc.save(path)
    patch_ooxml(path)


def write_catalog() -> None:
    catalog = {
        "version": 1,
        "templates": [
            {
                "id": "official-document",
                "path": "assets/templates/official-document.docx",
                "name": "正式公文模板",
                "use_for": ["决定", "通知", "通报", "请示", "报告", "批复", "意见"],
                "notes": "含版头、发文字号、红色分隔线、落款、版记占位。",
                "placeholders": [
                    "份号",
                    "密级和保密期限",
                    "紧急程度",
                    "发文机关标志",
                    "发文字号",
                    "标题",
                    "主送机关",
                    "正文第一段",
                    "一级标题",
                    "正文第二段",
                    "二级标题",
                    "正文第三段",
                    "附件名称",
                    "发文机关署名",
                    "成文日期",
                    "抄送机关",
                    "印发机关",
                    "印发日期",
                ],
            },
            {
                "id": "official-letter",
                "path": "assets/templates/official-letter.docx",
                "name": "函件与普通通知模板",
                "use_for": ["函", "会议通知", "邀请函", "一般事务通知"],
                "notes": "适合无红头文件要求的正式联系类文稿。",
                "placeholders": [
                    "发文机关标志",
                    "标题",
                    "主送机关",
                    "正文第一段",
                    "一级标题",
                    "正文第二段",
                    "二级标题",
                    "正文第三段",
                    "联系人",
                    "联系电话",
                    "发文机关署名",
                    "成文日期",
                ],
            },
            {
                "id": "meeting-material",
                "path": "assets/templates/meeting-material.docx",
                "name": "会议材料模板",
                "use_for": ["讲话稿", "汇报材料", "工作总结", "工作方案", "调研报告"],
                "notes": "适合领导讲话、汇报、方案、总结等会议材料。",
                "placeholders": [
                    "会议名称",
                    "标题",
                    "讲话人",
                    "日期",
                    "开场段",
                    "第一部分标题",
                    "第一部分正文",
                    "第二部分标题",
                    "第二部分正文",
                    "第三部分标题",
                    "第三部分正文",
                    "结束段",
                ],
            },
        ],
    }
    (TEMPLATE_DIR / "template-catalog.json").write_text(
        json.dumps(catalog, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main() -> None:
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    official_document(TEMPLATE_DIR / "official-document.docx")
    official_letter(TEMPLATE_DIR / "official-letter.docx")
    meeting_material(TEMPLATE_DIR / "meeting-material.docx")
    write_catalog()


if __name__ == "__main__":
    main()
