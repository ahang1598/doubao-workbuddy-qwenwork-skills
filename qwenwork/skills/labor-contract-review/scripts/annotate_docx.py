#!/usr/bin/env python3
from __future__ import annotations

import argparse
from datetime import datetime
from dataclasses import dataclass
import json
from pathlib import Path
import re
import sys

from docx import Document


try:
    from scripts._bootstrap import ensure_skill_import_root
except ModuleNotFoundError:
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from _bootstrap import ensure_skill_import_root

ROOT = ensure_skill_import_root(__file__)

from labor_contract_review.annotation_anchor import Paragraph, resolve_annotation_paragraph_id
from labor_contract_review.docx_comments import CommentExportResult, CommentPayload, export_docx_with_comments
from labor_contract_review.review_result import normalize_issue_for_annotation


@dataclass
class ReviewIssue:
    rule_id: str
    clause_title: str
    clause_position: str
    paragraph_id: str | None
    risk_level: str
    risk_type: str
    original_excerpt: str
    risk_message: str
    risk_analysis: str
    revision_suggestion: str
    comment_title: str
    comment_author: str
    revised_text_example: str = ""


def main() -> int:
    parser = argparse.ArgumentParser(description="根据 issues.json 在原 docx 上生成 Word 批注版。")
    parser.add_argument("--document", "-d", type=Path, required=True, help="原始 docx 文件路径。")
    parser.add_argument("--issues-json", "-i", type=Path, required=True, help="render_review_outputs.py 生成的 issues JSON。")
    parser.add_argument("--output", "-o", type=Path, default=None, help="批注版 docx 输出路径。")
    args = parser.parse_args()

    document_path = resolve_document_path(args.document)
    output_path = args.output or build_default_output_path(document_path)
    result, warnings = annotate_docx(document_path, args.issues_json, output_path)
    if result.status == "deliverable":
        print(f"批注合同下载地址：{result.output_path}")
        print(f"成功写入批注：{result.applied_count} 条；跳过：{result.skipped_count + len(warnings)} 条")
    elif result.status == "no_annotation_needed":
        print(result.message or "合同无风险，无需生成批注版。")
    else:
        print(result.message or "批注版生成失败。", file=sys.stderr)
    for warning in warnings:
        print(f"警告：{warning}", file=sys.stderr)
    return 0 if result.status != "failed" else 1


def annotate_docx(document_path: Path, issues_json_path: Path, output_path: Path):
    paragraphs = parse_docx_paragraphs(document_path)
    paragraph_id_map = {paragraph.paragraph_id: paragraph for paragraph in paragraphs}
    raw = json.loads(issues_json_path.read_text(encoding="utf-8"))
    issues_payload = raw.get("issues", raw) if isinstance(raw, dict) else raw
    if not isinstance(issues_payload, list):
        raise ValueError("issues-json 必须是风险项数组，或包含 issues 数组的对象。")
    if not issues_payload:
        return (
            CommentExportResult(
                output_path=output_path,
                applied_count=0,
                skipped_count=0,
                status="no_annotation_needed",
                message="合同无风险，无需生成批注版。",
            ),
            [],
        )

    comments: list[CommentPayload] = []
    warnings: list[str] = []
    normalized_count = 0
    blank_template_skip_count = 0
    for index, item in enumerate(issues_payload, start=1):
        try:
            normalized = normalize_issue_for_annotation(item)
            normalized_count += 1
            issue = ReviewIssue(
                rule_id=normalized["rule_id"],
                clause_title=normalized["clause_title"],
                clause_position=normalized["clause_position"],
                paragraph_id=normalized.get("paragraph_id") or None,
                risk_level=normalized["risk_level"],
                risk_type=normalized["risk_type"],
                original_excerpt=normalized["original_excerpt"],
                risk_message=normalized["risk_message"],
                risk_analysis=normalized["risk_analysis"],
                revision_suggestion=normalized["revision_suggestion"],
                revised_text_example=normalized["revised_text_example"],
                comment_title=normalized["comment_title"],
                comment_author=normalized["comment_author"],
            )
        except ValueError as exc:
            if "空白模板待填栏位" in str(exc):
                blank_template_skip_count += 1
            warnings.append(f"第 {index} 条风险项未生成批注：{exc}")
            continue

        paragraph_id = resolve_annotation_paragraph_id(
            paragraph_id=issue.paragraph_id,
            original_excerpt=issue.original_excerpt,
            clause_position=issue.clause_position,
            paragraph_id_map=paragraph_id_map,
            additional_clues=[
                issue.clause_title,
                issue.risk_message,
                issue.risk_analysis,
            ],
        )
        if not paragraph_id:
            warnings.append(
                f"第 {index} 条风险项未生成批注：未定位到可挂载段落"
                f"（{issue.clause_title}；{issue.clause_position}）。"
            )
            continue
        comments.append(
            CommentPayload(
                paragraph_id=paragraph_id,
                excerpt=issue.original_excerpt,
                body=comment_body(issue),
                author=issue.comment_author,
                comment_title=issue.comment_title,
            )
        )

    if not comments:
        if normalized_count == 0:
            if blank_template_skip_count == len(issues_payload):
                return (
                    CommentExportResult(
                        output_path=output_path,
                        applied_count=0,
                        skipped_count=len(issues_payload),
                        status="no_annotation_needed",
                        message="空白模板待填事项无需生成批注版。",
                    ),
                    warnings,
                )
            return (
                CommentExportResult(
                    output_path=output_path,
                    applied_count=0,
                    skipped_count=len(issues_payload),
                    status="failed",
                    message="批注版生成失败：存在命中风险，但风险项缺少生成批注所需字段。",
                ),
                warnings,
            )
        return (
            CommentExportResult(
                output_path=output_path,
                applied_count=0,
                skipped_count=normalized_count,
                status="failed",
                message="批注版生成失败：存在命中风险，但未定位到任何可挂载段落。",
            ),
            warnings,
        )

    return export_docx_with_comments(document_path, output_path, comments), warnings


def build_default_output_path(document_path: Path) -> Path:
    date_suffix = datetime.now().strftime("%Y%m%d")
    filename = f"{document_path.stem}-批注版-{date_suffix}.docx"
    return document_path.with_name(filename)


def resolve_document_path(raw_path: Path) -> Path:
    candidate = raw_path.expanduser()
    if candidate.exists():
        return candidate.resolve()

    name = candidate.name
    search_roots = [
        Path.cwd(),
        ROOT,
        ROOT / "output",
        Path("/tmp"),
    ]

    for root in search_roots:
        direct = root / name
        if direct.exists():
            return direct.resolve()

    for root in search_roots:
        try:
            for found in root.rglob(name):
                if found.is_file():
                    return found.resolve()
        except (OSError, PermissionError):
            continue

    raise FileNotFoundError(f"未找到原始 docx 文件: {raw_path}")


def parse_docx_paragraphs(path: Path) -> list[Paragraph]:
    if path.suffix.lower() != ".docx":
        raise ValueError("本地批注仅支持 docx 原件。")

    document = Document(path)
    paragraphs: list[Paragraph] = []

    for index, paragraph in enumerate(document.paragraphs, start=1):
        text = paragraph.text.strip()
        if not text:
            continue
        paragraphs.append(Paragraph(paragraph_id=f"p{index}", text=text, source="docx", location_hint=f"段落 {index}"))

    for t_index, table in enumerate(document.tables, start=1):
        for r_index, row in enumerate(table.rows, start=1):
            row_segments: list[str] = []
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            for c_index, cell in enumerate(row.cells, start=1):
                for p_index, paragraph in enumerate(cell.paragraphs, start=1):
                    text = paragraph.text.strip()
                    if not text:
                        continue
                    row_segments.append(text)
                    paragraphs.append(
                        Paragraph(
                            paragraph_id=f"t{t_index}r{r_index}c{c_index}p{p_index}",
                            text=text,
                            source="docx-table-cell",
                            location_hint=f"表格 {t_index} 行 {r_index} 列 {c_index} 段落 {p_index}",
                        )
                    )
            if not row_text and not row_segments:
                continue
            paragraphs.append(
                Paragraph(
                    paragraph_id=f"t{t_index}r{r_index}",
                    text=row_text or " | ".join(row_segments),
                    source="docx-table-row",
                    location_hint=f"表格 {t_index} 行 {r_index}",
                )
            )

    return paragraphs


def comment_body(issue: ReviewIssue) -> str:
    risk_note = to_plain_comment_text(build_combined_risk_note(issue.risk_message, issue.risk_analysis))
    revision_suggestion = to_plain_comment_text(issue.revision_suggestion)
    revised_text_example = to_plain_comment_text(issue.revised_text_example)
    sections = [
        f"【风险等级】{issue.risk_level}风险",
        f"【风险说明】{risk_note}",
        f"【修改建议】{revision_suggestion}",
    ]
    if revised_text_example:
        sections.extend(
            [
                "【修改后文本示例】",
                revised_text_example,
            ]
        )
    return "\n".join(sections)


def build_combined_risk_note(risk_message: str, risk_analysis: str) -> str:
    message = (risk_message or "").strip()
    analysis = (risk_analysis or "").strip()
    if not message:
        return analysis
    if not analysis:
        return message
    if message == analysis:
        return message
    return f"{message} {analysis}"


def to_plain_comment_text(text: str) -> str:
    value = (text or "").strip()
    if not value:
        return ""

    value = value.replace("\\r\\n", "\n").replace("\\n", "\n").replace("\\r", "\n")
    value = value.replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
    value = re.sub(r"```[\s\S]*?```", _strip_fenced_code, value)
    value = re.sub(r"`([^`]*)`", r"\1", value)

    cleaned_lines: list[str] = []
    for raw_line in value.splitlines():
        line = raw_line.strip()
        if not line:
            if cleaned_lines and cleaned_lines[-1] != "":
                cleaned_lines.append("")
            continue

        if re.fullmatch(r"\|?[\s:\-]+\|[\s\|:\-]*", line):
            continue

        if line.startswith("|") and line.endswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|") if cell.strip()]
            line = "；".join(cells)

        line = re.sub(r"^\s{0,3}#{1,6}\s*", "", line)
        line = re.sub(r"^\s*[-*+]\s+", "", line)
        line = re.sub(r"^\s*\d+[.)、]\s+", "", line)
        line = re.sub(r"(\*\*|__)(.*?)\1", r"\2", line)
        line = re.sub(r"(?<!\*)\*(?!\*)(.*?)\*(?<!\*)", r"\1", line)
        line = re.sub(r"(?<!_)_(?!_)(.*?)_(?<!_)", r"\1", line)
        line = re.sub(r"\[(.*?)\]\((.*?)\)", r"\1", line)
        line = re.sub(r"\[(.*?)\]", r"\1", line)
        line = line.replace("\\|", "|")
        line = re.sub(r"\s+", " ", line).strip()
        if line:
            cleaned_lines.append(line)

    while cleaned_lines and cleaned_lines[-1] == "":
        cleaned_lines.pop()
    return "\n".join(cleaned_lines)


def _strip_fenced_code(match: re.Match[str]) -> str:
    block = match.group(0)
    lines = block.splitlines()
    if len(lines) <= 2:
        return ""
    return "\n".join(lines[1:-1]).strip()


if __name__ == "__main__":
    raise SystemExit(main())
