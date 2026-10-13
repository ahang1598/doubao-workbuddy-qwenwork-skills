import json
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
import unittest

from docx import Document

from labor_contract_review.annotation_anchor import (
    Paragraph,
    match_paragraph_id_by_clause_position,
    match_paragraph_id_by_excerpt,
    resolve_annotation_paragraph_id,
)
from labor_contract_review.review_result import (
    ReportValidationError,
    extract_issues,
    load_raw_result,
    normalize_issue_for_annotation,
    validate_report_markdown,
    write_outputs,
)
from scripts.annotate_docx import ReviewIssue, annotate_docx
from scripts.annotate_report_docx import main as annotate_report_main


class ReviewResultTest(unittest.TestCase):
    def test_extract_issues_keeps_distinct_risks_at_same_position(self) -> None:
        raw = {
            "issues": [
                {
                    "rule_id": "LCR010",
                    "clause_title": "违约责任",
                    "clause_position": "第十条",
                    "original_excerpt": "甲方可对乙方罚款",
                    "risk_message": "罚款风险",
                    "risk_analysis": "罚款缺少合法依据。",
                    "revision_suggestion": "删除罚款表述。",
                },
                {
                    "rule_id": "LCR010",
                    "clause_title": "违约责任",
                    "clause_position": "第十条",
                    "original_excerpt": "乙方应赔偿甲方全部损失",
                    "risk_message": "赔偿范围过宽",
                    "risk_analysis": "赔偿责任未限定实际损失。",
                    "revision_suggestion": "限定为依法可主张的实际损失。",
                },
            ]
        }

        issues = extract_issues(raw)

        self.assertEqual(len(issues), 2)
        self.assertEqual({issue["risk_message"] for issue in issues}, {"罚款风险", "赔偿范围过宽"})

    def test_markdown_fallback_issues_can_anchor_paraphrased_excerpts(self) -> None:
        raw = load_raw_result(Path("references/example-review.json"))
        report_markdown = raw["report_markdown"]
        issues = extract_issues({"report_markdown": report_markdown}, report_markdown)
        paragraph_id_map = {
            "p1": Paragraph(
                paragraph_id="p1",
                text="甲乙双方同意，社会保险由乙方自行缴纳，甲方按月向乙方支付社保补贴。",
                source="docx",
                location_hint="段落 1",
            ),
            "p2": Paragraph(
                paragraph_id="p2",
                text="甲方安排乙方从事运营相关工作，具体岗位可根据甲方经营需要随时调整。乙方工作地点为全国范围内甲方指定地点，甲方可根据项目需要安排乙方长期驻外，无需另行协商。",
                source="docx",
                location_hint="段落 2",
            ),
            "p3": Paragraph(
                paragraph_id="p3",
                text="乙方如违反公司制度，甲方可视情况予以罚款，或直接从乙方工资中扣除相应金额。",
                source="docx",
                location_hint="段落 3",
            ),
        }

        target_issues = [
            issue for issue in issues if issue.get("rule_id") in {"LCR003", "LCR004", "LCR010"}
        ]
        matched = [
            match_paragraph_id_by_excerpt(issue["original_excerpt"], paragraph_id_map)
            for issue in target_issues
        ]

        self.assertEqual(matched, ["p2", "p1", "p3"])

    def test_example_review_preserves_hfr_issues_in_outputs(self) -> None:
        raw = load_raw_result(Path("references/example-review.json"))
        with tempfile.TemporaryDirectory() as tmpdir:
            report_path, issues_path = write_outputs(raw, Path(tmpdir), "review")

            report_markdown = report_path.read_text(encoding="utf-8")
            rendered_issues = json.loads(issues_path.read_text(encoding="utf-8"))["issues"]

        hfr_ids = {"HFR003", "HFR007", "HFR010", "HFR011"}

        for rule_id in hfr_ids:
            self.assertIn(rule_id, report_markdown)
        self.assertTrue(hfr_ids.issubset({issue["rule_id"] for issue in rendered_issues}))

    def test_attachment_section_allows_generated_docx_markdown_link(self) -> None:
        raw = load_raw_result(Path("references/example-review.json"))
        report_markdown = raw["report_markdown"].replace(
            "| 未提供 | 未提供 |",
            "| [门店店长-批注版-20260604.docx](门店店长-批注版-20260604.docx) | 带批注的合同版本 |",
        )

        validate_report_markdown(report_markdown)

    def test_example_review_issues_are_ordered_by_contract_clause(self) -> None:
        raw = load_raw_result(Path("references/example-review.json"))
        issues = extract_issues(raw, raw["report_markdown"])

        self.assertEqual(
            [issue["rule_id"] for issue in issues],
            ["LCR004", "LCR003", "LCR010", "HFR003", "HFR007", "HFR010", "HFR011"],
        )

    def test_healthcare_hfr_rules_are_documented_for_release(self) -> None:
        skill = Path("SKILL.md").read_text(encoding="utf-8")
        template = Path("references/report-template.md").read_text(encoding="utf-8")
        rules = Path("references/labor-contract-rules.md").read_text(encoding="utf-8")
        boundary = Path("references/boundary-rules.md").read_text(encoding="utf-8")
        output_contract = Path("references/output-contract.md").read_text(encoding="utf-8")

        for text in (skill, template):
            self.assertIn("HFR001-HFR015", text)
        for rule_id in ("HFR012", "HFR013", "HFR014", "HFR015"):
            self.assertIn(rule_id, rules)
        for keyword in ("持证", "驾驶证", "特种作业证", "轮班", "夜班", "社保缴纳起算", "暂停上岗"):
            self.assertIn(keyword, skill + rules + boundary + output_contract)

    def test_over_inference_guardrails_are_documented_for_release(self) -> None:
        skill = Path("SKILL.md").read_text(encoding="utf-8")
        template = Path("references/report-template.md").read_text(encoding="utf-8")
        rules = Path("references/labor-contract-rules.md").read_text(encoding="utf-8")
        boundary = Path("references/boundary-rules.md").read_text(encoding="utf-8")
        output_contract = Path("references/output-contract.md").read_text(encoding="utf-8")
        combined = skill + template + rules + boundary + output_contract

        for keyword in (
            "不得做过度推断",
            "原文没有",
            "保密条款不得扩展成薪资保密",
            "离职后继续保密义务",
            "标准工时条款不得扩展成加班审批例外",
            "入职后30日内办理参保手续",
            "表述歧义",
            "原则性依法赔偿条款不得扩展成解除或扣款风险",
            "不得写死固定补正天数",
            "不得设置“通勤增加30分钟”等固定阈值",
            "默认按中风险处理",
            "必须直接输出 `HFR003` 风险卡片",
            "由双方另行协商确定",
        ):
            self.assertIn(keyword, combined)

    def test_employer_side_guardrails_are_documented_for_release(self) -> None:
        skill = Path("SKILL.md").read_text(encoding="utf-8")
        template = Path("references/report-template.md").read_text(encoding="utf-8")
        rules = Path("references/labor-contract-rules.md").read_text(encoding="utf-8")
        boundary = Path("references/boundary-rules.md").read_text(encoding="utf-8")
        output_contract = Path("references/output-contract.md").read_text(encoding="utf-8")
        combined = skill + template + rules + boundary + output_contract

        for keyword in (
            "合法底线 + 企业保护",
            "默认按甲方/用人单位立场审查",
            "不得为了保护乙方而扩大甲方义务",
            "不得把法律没有强制要求的补贴",
            "管理弹性",
            "成本可控",
            "按甲方依法制定并公示的制度执行",
            "甲方合同硬义务",
        ):
            self.assertIn(keyword, combined)

    def test_skill_entrypoint_stays_small_and_references_detailed_contracts(self) -> None:
        skill = Path("SKILL.md").read_text(encoding="utf-8")

        self.assertLess(len(skill.splitlines()), 150)
        self.assertTrue(Path("references/boundary-rules.md").exists())
        self.assertTrue(Path("references/output-contract.md").exists())
        for keyword in (
            "目录说明",
            "references/boundary-rules.md",
            "references/output-contract.md",
            "# 劳动合同审查报告",
            "完整正式报告正文",
            "不能只输出摘要、交付成果或文件链接",
            "不得根据预期命名、报告文字或计划结果虚构 Word 附件",
            "异常处理与状态管理",
            "本技能为无状态设计",
            "scripts/",
            "labor_contract_review/",
            "agents/openai.yaml",
        ):
            self.assertIn(keyword, skill)

    def test_scripts_can_import_packaged_module_from_named_skill_directory(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir)
            skill_root = tmp_path / "labor-contract-review"
            for name in ("SKILL.md", "pyproject.toml", "agents", "labor_contract_review", "references", "scripts"):
                source = Path(name)
                target = skill_root / name
                if source.is_dir():
                    shutil.copytree(source, target)
                else:
                    target.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copy2(source, target)

            output_dir = tmp_path / "out"
            result = subprocess.run(
                [
                    sys.executable,
                    str(skill_root / "scripts" / "render_review_outputs.py"),
                    "--input",
                    str(skill_root / "references" / "example-no-risk-review.json"),
                    "--output-dir",
                    str(output_dir),
                    "--no-print-report",
                ],
                cwd=tmp_path,
                text=True,
                capture_output=True,
                check=False,
            )

            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue(any(output_dir.glob("*.md")))

    def test_risk_level_grouped_detail_headings_are_rejected(self) -> None:
        raw = load_raw_result(Path("references/example-review.json"))
        report_markdown = raw["report_markdown"].replace(
            "### 2. 按合同条款顺序的风险问题",
            "### 2. 高风险问题",
            1,
        )

        with self.assertRaises(ReportValidationError):
            validate_report_markdown(report_markdown)

    def test_risk_cards_must_use_project_content_header(self) -> None:
        raw = load_raw_result(Path("references/example-review.json"))
        report_markdown = raw["report_markdown"].replace(
            "#### 1. 工作内容和工作地点（规则编号：LCR004）\n\n| 项目 | 内容 |",
            "#### 1. 工作内容和工作地点（规则编号：LCR004）\n\n| 审查项 | 注意事项 |",
            2,
        )

        with self.assertRaises(ReportValidationError):
            validate_report_markdown(report_markdown)

    def test_missing_risk_detail_section_is_rejected_when_report_declares_risks(self) -> None:
        raw = load_raw_result(Path("references/example-review.json"))
        report_markdown = raw["report_markdown"]
        start = report_markdown.index("### 2. 按合同条款顺序的风险问题")
        end = report_markdown.index("### 6. 需用户确认或文本无法确定事项")
        report_markdown = report_markdown[:start] + report_markdown[end:]

        with self.assertRaises(ReportValidationError):
            validate_report_markdown(report_markdown)

    def test_no_risk_report_can_omit_risk_problem_section(self) -> None:
        raw = load_raw_result(Path("references/example-no-risk-review.json"))

        validate_report_markdown(raw["report_markdown"])

    def test_hfr_anchor_matches_source_paragraphs(self) -> None:
        paragraph_id_map = {
            "p7": Paragraph(
                "p7",
                "乙方未办理离职交接或交接不完整造成损失不能确定的，赔偿标准至少为月平均劳动报酬之2倍。",
                "docx",
                "段落 7",
            ),
            "p10": Paragraph(
                "p10",
                "乙方收到工资后7个工作日内未书面提出异议的，视为认可当月工资、加班费、提成等全部结算。",
                "docx",
                "段落 10",
            ),
            "p11": Paragraph(
                "p11",
                "双方发生劳动争议的，可向杭州市高新开发区（滨江）劳动人事争议调解委员会申请仲裁。",
                "docx",
                "段落 11",
            ),
        }
        cases = [
            ("HFR007", "乙方未办理离职交接或交接不完整造成损失不能确定的，赔偿标准至少为月平均劳动报酬之2倍。", "离职交接条款", "p7"),
            ("HFR010", "乙方收到工资后7个工作日内未书面提出异议的，视为认可当月工资、加班费、提成等全部结算。", "工资结算条款", "p10"),
            ("HFR011", "双方发生劳动争议的，可向杭州市高新开发区（滨江）劳动人事争议调解委员会申请仲裁。", "争议解决条款", "p11"),
        ]

        for rule_id, excerpt, position, expected_id in cases:
            with self.subTest(rule_id=rule_id):
                matched = resolve_annotation_paragraph_id(
                    paragraph_id=None,
                    original_excerpt=excerpt,
                    clause_position=position,
                    paragraph_id_map=paragraph_id_map,
                    additional_clues=[rule_id],
                )
                self.assertEqual(matched, expected_id)

    def test_hfr_and_lcr_same_position_are_not_deduped(self) -> None:
        raw = {
            "issues": [
                {
                    "rule_id": "LCR010",
                    "clause_title": "违约责任",
                    "clause_position": "第十条",
                    "original_excerpt": "甲方可直接从乙方工资中扣除相应金额。",
                    "risk_message": "工资扣罚风险",
                    "risk_analysis": "直接扣工资缺少合法依据和程序保障。",
                    "revision_suggestion": "删除直接扣工资表述。",
                },
                {
                    "rule_id": "HFR010",
                    "clause_title": "工资异议期限",
                    "clause_position": "第十条",
                    "original_excerpt": "乙方收到工资后7个工作日内未提出异议的，视为认可。",
                    "risk_message": "工资异议期限风险",
                    "risk_analysis": "内部核对期限不能剥夺依法追索工资差额的权利。",
                    "revision_suggestion": "明确该期限仅为内部核对期限。",
                },
            ]
        }

        issues = extract_issues(raw)

        self.assertEqual(len(issues), 2)
        self.assertEqual({issue["rule_id"] for issue in issues}, {"LCR010", "HFR010"})

    def test_blank_template_fields_are_not_extracted_as_issues(self) -> None:
        raw = {
            "issues": [
                {
                    "rule_id": "LCR006",
                    "clause_title": "劳动报酬",
                    "clause_position": "劳动报酬条款",
                    "risk_level": "高",
                    "risk_type": "工资金额缺失风险",
                    "original_excerpt": "乙方月工资为人民币____元，发薪日期为每月____日。",
                    "risk_message": "工资和发薪日期为空白待填。",
                    "risk_analysis": "该合同为打印空白模板，签署前应补齐工资和发薪日期。",
                    "revision_suggestion": "签署前填写并复核最终版本。",
                }
            ]
        }

        self.assertEqual(extract_issues(raw), [])

    def test_blank_subject_template_fields_are_not_extracted_as_issues(self) -> None:
        raw = {
            "issues": [
                {
                    "rule_id": "LCR001",
                    "clause_title": "乙方主体信息缺失",
                    "clause_position": "合同首部乙方信息栏",
                    "risk_level": "高",
                    "risk_type": "主体信息缺失",
                    "original_excerpt": "乙方（劳动者）：\\n身份证号码：____________\\n地址：____________\\n联系电话：____________",
                    "risk_message": "乙方信息缺失将导致合同主体不明确。",
                    "risk_analysis": "该合同为打印空白模板，相关栏位属于签署前待填写事项。",
                    "revision_suggestion": "签署前补齐并核验乙方姓名、身份证号码、地址和联系电话。",
                }
            ]
        }

        self.assertEqual(extract_issues(raw), [])

    def test_overtime_fixed_base_issue_keeps_practical_revision_boundary(self) -> None:
        raw = {
            "issues": [
                {
                    "rule_id": "LCR006",
                    "clause_title": "加班工资",
                    "clause_position": "劳动报酬条款",
                    "risk_level": "中",
                    "risk_type": "加班工资基数固化风险",
                    "original_excerpt": "加班需经甲方书面审批，加班工资以乙方月工资9500元为基数，按法定标准计算。",
                    "risk_message": "加班工资基数直接固定为月工资9500元，可能固化计算标准。",
                    "risk_analysis": "书面审批可以作为管理规则，但不能绝对排除实际安排或确认的加班。",
                    "revision_suggestion": "改为加班审批和依法支付原则，具体基数按依法制定并公示的薪酬、考勤和加班审批制度执行；保留甲方实际安排、确认或依法认定加班的例外。",
                }
            ]
        }

        issues = extract_issues(raw)

        self.assertEqual(len(issues), 1)
        self.assertIn("依法制定并公示", issues[0]["revision_suggestion"])
        self.assertNotIn("以月工资9500元为基数", issues[0]["revision_suggestion"])
        self.assertNotIn("一律不构成加班", issues[0]["revision_suggestion"])

    def test_comment_author_is_forced_to_contract_review_assistant(self) -> None:
        issue = normalize_issue_for_annotation(
            {
                "rule_id": "LCR017",
                "clause_title": "规章制度",
                "clause_position": "第十二条",
                "risk_level": "中",
                "risk_type": "规章制度公示风险",
                "original_excerpt": "合同多次提及乙方应遵守甲方规章制度。",
                "risk_message": "规章制度效力需要公示告知。",
                "risk_analysis": "未公示告知的规章制度难以作为处理依据。",
                "revision_suggestion": "补充规章制度已依法制定并公示告知。",
                "comment_author": "悟空AI助手",
            }
        )

        self.assertEqual(issue["comment_author"], "劳动合同审查助手")

    def test_annotate_docx_returns_no_annotation_needed_for_blank_template_only_issues(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir)
            document_path = tmp_path / "contract.docx"
            issues_path = tmp_path / "issues.json"
            output_path = tmp_path / "contract-批注版.docx"
            document = Document()
            document.add_paragraph("乙方（劳动者）：身份证号码：____________ 地址：____________ 联系电话：____________")
            document.save(document_path)
            issues_path.write_text(
                json.dumps(
                    {
                        "issues": [
                            {
                                "rule_id": "LCR001",
                                "clause_title": "乙方主体信息缺失",
                                "clause_position": "合同首部乙方信息栏",
                                "risk_level": "高",
                                "risk_type": "主体信息缺失",
                                "original_excerpt": "乙方（劳动者）：身份证号码：____________ 地址：____________ 联系电话：____________",
                                "risk_message": "乙方信息空白待填。",
                                "risk_analysis": "该合同为打印空白模板，相关栏位属于签署前待填写事项。",
                                "revision_suggestion": "签署前补齐并核验乙方姓名、身份证号码、地址和联系电话。",
                            }
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result, warnings = annotate_docx(document_path, issues_path, output_path)

        self.assertEqual(result.status, "no_annotation_needed")
        self.assertEqual(result.message, "空白模板待填事项无需生成批注版。")
        self.assertEqual(result.skipped_count, 1)
        self.assertEqual(len(warnings), 1)

    def test_annotate_docx_skips_blank_template_issue_and_keeps_real_risk(self) -> None:
        real_excerpt = "双方发生劳动争议的，可向杭州市高新开发区（滨江）劳动人事争议调解委员会申请仲裁。"
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir)
            document_path = tmp_path / "contract.docx"
            issues_path = tmp_path / "issues.json"
            output_path = tmp_path / "contract-批注版.docx"
            document = Document()
            document.add_paragraph("乙方（劳动者）：身份证号码：____________ 地址：____________ 联系电话：____________")
            document.add_paragraph(real_excerpt)
            document.save(document_path)
            issues_path.write_text(
                json.dumps(
                    {
                        "issues": [
                            {
                                "rule_id": "LCR001",
                                "clause_title": "乙方主体信息缺失",
                                "clause_position": "合同首部乙方信息栏",
                                "risk_level": "高",
                                "risk_type": "主体信息缺失",
                                "original_excerpt": "乙方（劳动者）：身份证号码：____________ 地址：____________ 联系电话：____________",
                                "risk_message": "乙方信息空白待填。",
                                "risk_analysis": "该合同为打印空白模板，相关栏位属于签署前待填写事项。",
                                "revision_suggestion": "签署前补齐并核验乙方姓名、身份证号码、地址和联系电话。",
                            },
                            {
                                "rule_id": "HFR011",
                                "clause_title": "争议解决",
                                "clause_position": "争议解决条款",
                                "risk_level": "低",
                                "risk_type": "争议解决机构表述错误",
                                "original_excerpt": real_excerpt,
                                "risk_message": "调解委员会和仲裁程序混用。",
                                "risk_analysis": "劳动争议仲裁应向有管辖权的劳动人事争议仲裁委员会申请。",
                                "revision_suggestion": "改为依法向有管辖权的劳动人事争议仲裁委员会申请仲裁。",
                            },
                        ]
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )

            result, warnings = annotate_docx(document_path, issues_path, output_path)
            output_exists = output_path.exists()

        self.assertEqual(result.status, "deliverable")
        self.assertEqual(result.applied_count, 1)
        self.assertTrue(output_exists)
        self.assertEqual(len(warnings), 1)

    def test_annotate_report_docx_generates_comments_from_markdown_report(self) -> None:
        report_markdown = """# 劳动合同审查报告

## 一、核心结论

| 项目 | 内容 |
| --- | --- |
| 整体风险等级 | 🔴 高风险 |
| 必须修改（不改会出事） | 1. 争议解决机构错误。 |
| 建议修改（优化项） | 1. 修正争议解决条款。 |
| 一句话总结 | 争议解决机构需要修正。 |

## 二、关键风险（精简版）

| 风险等级 | 风险摘要 |
| --- | --- |
| 🔴 高风险 | 争议解决机构错误。 |

## 三、针对行业的特别说明（自动生成）

| 项目 | 内容 |
| --- | --- |
| 行业类型 | 未明确 |
| 行业说明 | 通用劳动合同审查。 |

## 四、完整审查明细

### 2. 按合同条款顺序的风险问题

#### 1. 争议解决（规则编号：HFR011）

| 项目 | 内容 |
| --- | --- |
| 风险等级 | 🔴 高风险 |
| 风险类型 | 争议解决机构表述错误 |
| 原文位置及摘要 | 争议解决条款：双方发生劳动争议的，可向杭州市高新开发区（滨江）劳动人事争议调解委员会申请仲裁。 |
| 风险说明 | 调解委员会和仲裁程序混用。劳动争议仲裁应向有管辖权的劳动人事争议仲裁委员会申请。 |
| 修改建议 | 改为依法向有管辖权的劳动人事争议仲裁委员会申请仲裁。 |
| 修改后文本示例 | 双方发生劳动争议的，可依法向有管辖权的劳动人事争议仲裁委员会申请仲裁。 |

### 6. 需用户确认或文本无法确定事项

| 审查项 | 注意事项 |
| --- | --- |
| 无 | 无。 |

### 7. 整体修改建议

| 项目 | 内容 |
| --- | --- |
| 建议 1 | 修正争议解决机构。 |

### 8. 谈判与签署建议

1. 签署前修正争议解决机构。
2. 保留最终签署版本。

## 五、其他合规提示

1. 无。

## 六、附件下载

| 附件名称 | 说明 |
| --- | --- |
| 未提供 | 未提供 |
"""
        real_excerpt = "双方发生劳动争议的，可向杭州市高新开发区（滨江）劳动人事争议调解委员会申请仲裁。"
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp_path = Path(tmpdir)
            report_path = tmp_path / "report.md"
            document_path = tmp_path / "contract.docx"
            output_path = tmp_path / "contract-批注版.docx"
            report_path.write_text(report_markdown, encoding="utf-8")
            document = Document()
            document.add_paragraph(real_excerpt)
            document.save(document_path)

            old_argv = sys.argv
            try:
                sys.argv = [
                    "annotate_report_docx.py",
                    "--report",
                    str(report_path),
                    "--document",
                    str(document_path),
                    "--output",
                    str(output_path),
                ]
                exit_code = annotate_report_main()
            finally:
                sys.argv = old_argv

            self.assertEqual(exit_code, 0)
            self.assertTrue(output_path.exists())

    def test_blank_template_fields_cannot_be_reported_as_risk_cards(self) -> None:
        raw = load_raw_result(Path("references/example-review.json"))
        report_markdown = raw["report_markdown"].replace(
            "#### 1. 社会保险（规则编号：LCR003）",
            "#### 1. 工资空白待填（规则编号：LCR006）",
            1,
        ).replace(
            "社会保险条款：甲乙双方同意，社会保险由乙方自行缴纳，甲方按月向乙方支付社保补贴。",
            "劳动报酬条款：乙方月工资为人民币____元，发薪日期为每月____日。",
            1,
        ).replace(
            "社会保险缴纳义务原则上由用人单位依法办理并代扣代缴个人承担部分。将法定参保义务整体转给劳动者，争议时通常难以得到支持。",
            "工资和发薪日期为空白待填，应在签署前补齐并复核最终版本。",
            1,
        )

        with self.assertRaises(ReportValidationError):
            validate_report_markdown(report_markdown)

    def test_template_management_advice_is_not_extracted_as_issues(self) -> None:
        report_markdown = """# 劳动合同审查报告

## 四、完整审查明细

### 7. 整体修改建议

| 审查项 | 注意事项 |
| --- | --- |
| 建议 1 | 按普通员工、项目制员工、核心技术/管理人员区分模板。 |
| 建议 2 | 建立固定期限续签、连续两次固定期限和满十年等期限台账。 |

### 8. 谈判与签署建议

| 审查项 | 注意事项 |
| --- | --- |
| 签署动作 | 补齐附件签收、竞业启动/解除机制和送达地址留痕。 |
"""

        issues = extract_issues({"report_markdown": report_markdown}, report_markdown)

        self.assertEqual(issues, [])

    def test_write_outputs_removes_html_breaks_from_report_markdown(self) -> None:
        raw = load_raw_result(Path("references/example-review.json"))
        raw["report_markdown"] = raw["report_markdown"].replace(
            "1. 社会保险责任转嫁给劳动者。 2. 工作地点约定为全国范围且允许长期驻外。",
            "1. 社会保险责任转嫁给劳动者。<br>2. 工作地点约定为全国范围且允许长期驻外。\\n3. 规章制度未告知仍当然有效。",
        )
        report_path, _ = write_outputs(raw, Path("/tmp/lcr-test-no-html"), "review")

        report_markdown = report_path.read_text(encoding="utf-8")

        self.assertNotIn("<br>", report_markdown)
        self.assertNotIn("<br/>", report_markdown)
        self.assertNotIn("\\n", report_markdown)
        self.assertIn("1. 社会保险责任转嫁给劳动者。 2. 工作地点约定为全国范围且允许长期驻外。 3. 规章制度未告知仍当然有效。", report_markdown)

    def test_resolve_anchor_uses_additional_clues_when_excerpt_is_abstract(self) -> None:
        paragraph_id_map = {
            "p1": Paragraph("p1", "甲方无统一社会信用代码，乙方无手机、邮箱、紧急联系人。", "docx", "段落 1"),
            "p2": Paragraph("p2", "乙方工作地点为全国范围内甲方指定地点。", "docx", "段落 2"),
        }

        matched = resolve_annotation_paragraph_id(
            paragraph_id=None,
            original_excerpt="主体信息缺失",
            clause_position="甲乙方信息栏",
            paragraph_id_map=paragraph_id_map,
            additional_clues=["甲方无统一社会信用代码，乙方无手机/邮箱/紧急联系人"],
        )

        self.assertEqual(matched, "p1")

    def test_clause_position_can_match_article_heading_number(self) -> None:
        paragraph_id_map = {
            "p1": Paragraph("p1", "第六条 工作时间和休息休假", "docx", "段落 1"),
            "p2": Paragraph("p2", "第七条 劳动报酬 乙方离岗待岗期间工资按公司制度执行。", "docx", "段落 2"),
        }

        matched = match_paragraph_id_by_clause_position("第七条第三款", paragraph_id_map)

        self.assertEqual(matched, "p2")

    def test_annotation_anchor_does_not_use_revision_text_as_source_anchor(self) -> None:
        issue = ReviewIssue(
            rule_id="LCR010",
            clause_title="违反劳动合同责任",
            clause_position="第十条第6项",
            paragraph_id=None,
            risk_level="中",
            risk_type="工资扣罚风险",
            original_excerpt="工资扣除赔偿款",
            risk_message="甲方可直接从工资中扣除赔偿款，但未设置损失认定程序和每月扣除上限。",
            risk_analysis="工资扣除应以实际损失和法定扣除比例为基础。",
            revision_suggestion="删除第七条第三款第 7 项泄露/打听薪资条款。",
            revised_text_example="同时删除将打听薪资列为严重违纪的条款。",
            comment_title="违反劳动合同责任风险提示",
            comment_author="劳动合同审查助手",
        )
        paragraph_id_map = {
            "p6": Paragraph("p6", "6. 乙方给甲方造成经济损失的，甲方可要求赔偿并直接从乙方工资中扣除。", "docx", "段落 6"),
            "p7": Paragraph("p7", "7. 直接、间接泄露自己薪资，或者刺探、打听公司其他员工薪资，给公司造成不良影响的。", "docx", "段落 7"),
        }

        matched = resolve_annotation_paragraph_id(
            paragraph_id=issue.paragraph_id,
            original_excerpt=issue.original_excerpt,
            clause_position=issue.clause_position,
            paragraph_id_map=paragraph_id_map,
            additional_clues=[
                issue.clause_title,
                issue.risk_message,
                issue.risk_analysis,
            ],
        )

        self.assertEqual(matched, "p6")


if __name__ == "__main__":
    unittest.main()
