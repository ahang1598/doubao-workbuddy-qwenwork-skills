#!/usr/bin/env python3
"""
将预处理后的 Excel 数据导入多维表

读取合并后的 Excel 文件，按 schema 做字段类型转换，分批写入 aitable。

用法：
    python scripts/import_sku_records.py --base-id <BASE_ID> --table-id <TABLE_ID> --input "合并后.xlsx"

输出：JSON { success, total, imported, failed, errors }
"""

import json
import shutil
import subprocess
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print(json.dumps({"success": False, "error": "缺少依赖: openpyxl，请执行 pip install openpyxl"}))
    sys.exit(1)

SCHEMA_PATH = Path(__file__).resolve().parent.parent / "references" / "sku_table_field_schema.json"
BATCH_SIZE = 5

# 从 Excel 导入时跳过的字段（autoNumber 自动生成，attachment 需要手动上传）
SKIP_FIELD_TYPES = {"autoNumber", "attachment"}

# Excel 表头别名映射：当 Excel 表头与 schema 字段名不一致时，映射到 schema 字段名
# 例如商陆花的表头可能是"价格1"而非"价格"
FIELD_ALIAS_MAP = {
    "价格1": "价格",
}

# ── dws 工具函数 ──────────────────────────────────────────────────────────────

DWS_SEARCH_PATHS = [
    "~/.qoderwork/bin/dws",
    "~/.local/bin/dws",
    "dws",
]


def resolve_dws_binary() -> str:
    """查找 dws 可执行文件路径。"""
    for candidate in DWS_SEARCH_PATHS:
        expanded = str(Path(candidate).expanduser())
        if Path(expanded).is_file():
            return expanded
    found = shutil.which("dws")
    if found:
        return found
    print(json.dumps({"success": False, "error": "找不到 dws 命令，请确认已安装 dws CLI 工具"}))
    sys.exit(1)


def run_dws(args: list[str], timeout: int = 120) -> dict | None:
    """执行 dws 命令并返回解析后的 JSON 结果。"""
    dws_bin = resolve_dws_binary()
    cmd = [dws_bin] + args
    print(f"执行命令: {' '.join(cmd)}", file=sys.stderr)
    try:
        result = subprocess.run(cmd, capture_output=True, text=True, encoding="utf-8", timeout=timeout)
        if result.returncode != 0:
            print(f"错误：{result.stderr.strip()}", file=sys.stderr)
            return None
        return json.loads(result.stdout)
    except subprocess.TimeoutExpired:
        print(f"错误：命令执行超时（{timeout}秒）", file=sys.stderr)
        return None
    except json.JSONDecodeError:
        print(f"无法解析响应：{result.stdout[:300]}...", file=sys.stderr)
        return None
    except FileNotFoundError:
        print("错误：未找到 dws 命令", file=sys.stderr)
        return None


# ── 数据转换 ──────────────────────────────────────────────────────────────────

def load_schema() -> dict:
    """加载 schema，返回 {fieldName: {type, config}} 映射。"""
    with open(SCHEMA_PATH, encoding="utf-8") as schema_file:
        schema = json.load(schema_file)
    field_map = {}
    for field in schema["fields"]:
        field_map[field["fieldName"]] = field
    return field_map


def convert_cell_value(value, field_schema: dict):
    """根据字段 schema 转换单元格的值。"""
    field_type = field_schema["type"]

    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None

    if field_type == "number":
        try:
            num = float(value)
            formatter = field_schema.get("config", {}).get("formatter", "")
            if formatter == "INT":
                return int(num)
            return num
        except (ValueError, TypeError):
            return None

    if field_type == "date":
        return str(value) if value else None

    # text 及其他类型统一转字符串
    return str(value) if value is not None else None


def read_excel_records(input_path: str, field_map: dict) -> list[dict]:
    """读取 Excel，返回转换后的记录列表。"""
    workbook = openpyxl.load_workbook(input_path, read_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]

    # 确定需要导入的列（存在于 schema 中且不是跳过类型的字段）
    # 支持别名映射：Excel 表头通过 FIELD_ALIAS_MAP 映射到 schema 字段名
    importable_columns = []
    for col_index, header in enumerate(headers):
        mapped_name = FIELD_ALIAS_MAP.get(header, header)
        if mapped_name in field_map and field_map[mapped_name]["type"] not in SKIP_FIELD_TYPES:
            importable_columns.append((col_index, mapped_name))

    records = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        cells = {}
        for col_index, field_name in importable_columns:
            converted_value = convert_cell_value(row[col_index], field_map[field_name])
            if converted_value is not None:
                cells[field_name] = converted_value
        if cells:
            records.append(cells)

    workbook.close()
    return records


# ── 主逻辑 ────────────────────────────────────────────────────────────────────

def import_records(base_id: str, table_id: str, records: list[dict]) -> dict:
    """分批将记录写入多维表。"""
    result = {"success": True, "total": len(records), "imported": 0, "failed": 0, "errors": []}
    total = len(records)

    for i in range(0, total, BATCH_SIZE):
        batch = records[i:i + BATCH_SIZE]
        dws_records = [{"cells": rec} for rec in batch]
        records_json = json.dumps(dws_records, ensure_ascii=False)

        batch_end = min(i + BATCH_SIZE, total)
        print(f"正在写入第 {i + 1}-{batch_end} 条（共 {total} 条）...", file=sys.stderr)

        dws_result = run_dws([
            "aitable", "record", "create",
            "--base-id", base_id,
            "--table-id", table_id,
            "--records", records_json,
            "--format", "json",
            "-y",
        ])

        if dws_result:
            result["imported"] += len(batch)
            print(f"  ✅ 成功", file=sys.stderr)
        else:
            result["failed"] += len(batch)
            result["errors"].append(f"批次 {i + 1}-{batch_end} 写入失败")
            print(f"  ❌ 写入失败", file=sys.stderr)

    if result["failed"] > 0:
        result["success"] = False

    return result


def main():
    import argparse
    parser = argparse.ArgumentParser(description="将预处理后的 Excel 数据导入多维表")
    parser.add_argument("--base-id", required=True, help="多维表 Base ID")
    parser.add_argument("--table-id", required=True, help="多维表 Table ID")
    parser.add_argument("--input", required=True, help="预处理后的 Excel 文件路径")
    args = parser.parse_args()

    import os
    if not os.path.exists(args.input):
        print(json.dumps({"success": False, "error": f"文件不存在: {args.input}"}))
        sys.exit(1)

    field_map = load_schema()
    print(f"读取 Excel: {args.input}", file=sys.stderr)
    records = read_excel_records(args.input, field_map)
    print(f"共读取 {len(records)} 条记录", file=sys.stderr)

    if not records:
        print(json.dumps({"success": True, "total": 0, "imported": 0, "failed": 0, "errors": []}))
        return

    result = import_records(args.base_id, args.table_id, records)
    print(json.dumps(result, ensure_ascii=False))

    if not result["success"]:
        sys.exit(1)


if __name__ == "__main__":
    main()
