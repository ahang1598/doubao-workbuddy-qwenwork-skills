#!/usr/bin/env python3
"""
将 Excel 表格中相同款号的行合并：条码、颜色、尺码、库存 用逗号分隔，其他字段取第一行的值。

用法:
    python scripts/merge_sku.py --input "输入文件.xlsx" --output "输出文件.xlsx"

如果不指定 --output，默认在输入文件同目录下生成 `原文件名_合并.xlsx`。
"""

import argparse
import json
import os
import sys
from collections import OrderedDict

try:
    import openpyxl
except ImportError:
    print(json.dumps({"success": False, "error": "缺少依赖: openpyxl，请执行 pip install openpyxl"}))
    sys.exit(1)

MERGE_COLUMNS = ["条码", "颜色", "尺码", "库存"]
GROUP_BY_COLUMN = "款号"


def main():
    parser = argparse.ArgumentParser(description="合并相同款号的 SKU 行")
    parser.add_argument("--input", required=True, help="输入 Excel 文件路径")
    parser.add_argument("--output", default=None, help="输出 Excel 文件路径（可选）")
    args = parser.parse_args()

    input_path = args.input
    if not os.path.exists(input_path):
        print(json.dumps({"success": False, "error": f"文件不存在: {input_path}"}))
        sys.exit(1)

    # 默认输出路径
    if args.output:
        output_path = args.output
    else:
        base, ext = os.path.splitext(input_path)
        output_path = f"{base}_合并{ext}"

    workbook = openpyxl.load_workbook(input_path)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]

    # 校验必需列是否存在
    missing_columns = [col for col in [GROUP_BY_COLUMN] + MERGE_COLUMNS if col not in headers]
    if missing_columns:
        print(json.dumps({"success": False, "error": f"表头缺少必需列: {missing_columns}"}))
        sys.exit(1)

    merge_column_indices = [headers.index(col) for col in MERGE_COLUMNS]
    group_by_index = headers.index(GROUP_BY_COLUMN)

    # 按款号分组，保持原始顺序
    grouped_data = OrderedDict()
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, values_only=True):
        group_key = row[group_by_index]
        if group_key is None:
            continue
        if group_key not in grouped_data:
            grouped_data[group_key] = []
        grouped_data[group_key].append(list(row))

    # 创建输出工作簿
    output_workbook = openpyxl.Workbook()
    output_sheet = output_workbook.active
    output_sheet.title = sheet.title
    output_sheet.append(headers)

    merged_count = 0
    for group_key, rows in grouped_data.items():
        merged_row = list(rows[0])

        if len(rows) > 1:
            for col_index in merge_column_indices:
                values = [str(r[col_index]) if r[col_index] is not None else "" for r in rows]
                merged_row[col_index] = ",".join(values)
            merged_count += 1

        output_sheet.append(merged_row)

    output_workbook.save(output_path)

    result = {
        "success": True,
        "input_file": input_path,
        "output_file": output_path,
        "original_row_count": sheet.max_row - 1,
        "merged_row_count": len(grouped_data),
        "styles_with_multiple_skus": merged_count,
    }
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
