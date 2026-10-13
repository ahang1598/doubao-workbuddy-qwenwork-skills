#!/usr/bin/env python3
"""Generate golden_cases xlsx from current case definitions."""
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Golden Cases v4"

headers = ["评测任务ID", "技能名称", "技能英文名", "一级分类", "二级分类", "query", "检查点", "复杂度", "依赖case", "责任人"]
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border

rows = [
    # --- 信息查询类（低） ---
    ["case-02", "门店任务派发系统", "store-inspection", "门店巡检", "交互体验",
     "请使用 store-inspection 技能，介绍一下这个巡检技能都有哪些功能和能力？",
     "1. 助手成功介绍了技能的核心能力，包括任务派发、巡检报告生成等\n2. 介绍内容基于技能实际能力，未编造不存在的功能",
     "低", "", "葭牧"],
    ["case-04", "门店任务派发系统", "store-inspection", "门店巡检", "门店管理",
     "请使用 store-inspection 技能，帮我查看当前所有门店列表。",
     "1. 助手通过技能 CLI 的 store list 命令获取了门店列表\n2. 助手向用户展示了门店信息（如果系统中没有门店则正确提示为空）",
     "低", "", "葭牧"],
    ["case-06", "门店任务派发系统", "store-inspection", "门店巡检", "消息文案",
     "请使用 store-inspection 技能，帮我起草一份巡检任务的派发通知消息文案，语气要正式，包含检查项和截止时间提醒。假设检查项有「货架陈列」「收银台整洁度」「地面清洁」，截止时间明天18:00。",
     "1. 助手生成了包含检查项和截止时间信息的派发通知文案\n2. 文案语气正式，适合正式工作通知场景\n3. 助手向用户展示了文案内容",
     "低", "", "葭牧"],
    # --- 基础管理类（中） ---
    ["case-01", "门店任务派发系统", "store-inspection", "门店巡检", "环境初始化",
     "请使用 store-inspection 技能，帮我初始化一个全新的巡检环境。品牌叫「评测专用店」。请创建新的 AI 表格。",
     "1. 助手通过技能 CLI 成功执行了初始化操作（init create 或 init validate）\n2. 初始化结果中包含 baseId 等关键配置信息\n3. 助手向用户反馈了初始化完成的结果",
     "中", "", "葭牧"],
    ["case-03", "门店任务派发系统", "store-inspection", "门店巡检", "门店管理",
     "请使用 store-inspection 技能，帮我添加3家门店：评测A店（userId 440535）、评测B店（userId 151046）、评测C店（userId 427860）。直接使用这些 userId，不需要搜索通讯录。",
     "1. 助手通过技能 CLI 的 store add 命令添加了门店\n2. 每家门店正确关联了对应的店长 userId\n3. 添加完成后向用户展示了门店列表或确认结果",
     "中", "", "葭牧"],
    ["case-16", "门店任务派发系统", "store-inspection", "门店巡检", "配置分享",
     "请使用 store-inspection 技能，帮我生成一份巡检配置分享内容，我想分享给同事。先帮我生成分享内容就行，不用现在发送给任何人。",
     "1. 助手通过技能 CLI 的 share 命令生成了分享内容或向用户展示了当前巡检配置摘要\n2. 助手未擅自将分享消息发送给任何人",
     "中", "", "葭牧"],
    ["case-14", "门店任务派发系统", "store-inspection", "门店巡检", "统计分析",
     "请使用 store-inspection 技能，帮我查看各门店的巡检表现统计，看看哪家门店表现最好。",
     "1. 助手通过技能 CLI 的 stats 命令获取了统计数据，或正确告知用户当前没有足够的历史数据\n2. 助手向用户清晰展示了查询结果",
     "中", "", "葭牧"],
    ["case-15", "门店任务派发系统", "store-inspection", "门店巡检", "统计分析",
     "请使用 store-inspection 技能，帮我查看模板的使用率和任务完成率统计。",
     "1. 助手通过技能 CLI 的 stats 命令获取了模板统计数据，或正确告知用户当前没有足够的历史数据\n2. 助手向用户清晰展示了查询结果",
     "中", "", "葭牧"],
    ["case-09", "门店任务派发系统", "store-inspection", "门店巡检", "监控催办",
     "请使用 store-inspection 技能，帮我查看当前巡检任务的完成进度，有哪些门店已经提交了，哪些还没有？",
     "1. 助手通过技能 CLI 查询了巡检任务的完成进度\n2. 助手向用户清晰展示了查询结果（各门店提交状态或当前无可查询的任务）",
     "中", "", "葭牧"],
    # --- 有条件管理类（中） ---
    ["case-17", "门店任务派发系统", "store-inspection", "门店巡检", "门店管理",
     "请使用 store-inspection 技能，帮我查看当前门店列表，然后删除列表中最后一家门店。",
     "1. 助手先查询了当前门店列表\n2. 助手在执行删除操作前向用户确认了操作意图（删除是高危操作需确认）\n3. 助手执行了删除并向用户反馈了结果，或如果没有门店则正确告知用户",
     "中", "case-03（需要有门店可删）", "葭牧"],
    ["case-10", "门店任务派发系统", "store-inspection", "门店巡检", "监控催办",
     "请使用 store-inspection 技能，对还没提交巡检的门店发送催办提醒消息。",
     "1. 助手通过技能 CLI 查询了当前任务状态\n2. 助手根据查询结果采取正确行动：有未提交门店且有有效form_url则催办；全部完成则告知无需催办；无可用派发任务则正确告知用户——以上均视为通过\n3. 助手向用户反馈了最终结果，且未在无有效form_url时强行发送通知",
     "中", "case-07（需要有派发任务）", "葭牧"],
    # --- 模板与派发类（高） ---
    ["case-05", "门店任务派发系统", "store-inspection", "门店巡检", "模板管理",
     "请使用 store-inspection 技能，帮我创建一个巡检模板，叫「评测用巡检模板」，检查项有：货架陈列、收银台整洁度、地面清洁。每项需要拍照并AI审核。",
     "1. 助手通过技能 CLI 的 template create 命令创建了巡检模板\n2. 模板包含了用户指定的3个检查项\n3. 助手向用户反馈了创建结果",
     "高", "", "葭牧"],
    ["case-18", "门店任务派发系统", "store-inspection", "门店巡检", "模板管理",
     "请使用 store-inspection 技能，帮我先查看已有模板列表，然后选择一个模板进行修改：新增一个检查项「员工仪容仪表」，需要拍照和AI审核。",
     "1. 助手先通过 template list 查询了可用模板\n2. 助手通过技能 CLI 的 template update 命令执行了模板修改，或如果没有可用模板则正确告知用户\n3. 助手向用户反馈了操作结果",
     "高", "case-05（需要有模板可改）", "葭牧"],
    ["case-07", "门店任务派发系统", "store-inspection", "门店巡检", "任务派发",
     "请使用 store-inspection 技能，帮我派发一次巡检任务，截止时间设置为明天18:00。如果系统中还没有巡检模板，请先帮我创建一个叫「评测派发模板」的模板，检查项为：货架陈列、地面清洁，每项需要拍照和AI审核。然后再用该模板派发任务。",
     "1. 助手最终通过 dispatch init 创建了派发任务实例（无模板时先 template create 再 dispatch init 也算通过）\n2. 助手设置了截止时间或在 extra_info 中包含截止时间信息\n3. 助手反馈了派发结果；表单链接因同步延迟无法获取时正确引导用户完成发布也算通过",
     "高", "case-05（可选，prompt已自包含创建模板指令）", "葭牧"],
    ["case-08", "门店任务派发系统", "store-inspection", "门店巡检", "任务派发",
     "请使用 store-inspection 技能，帮我查看系统中已有的派发任务，然后给最近一次派发任务的所有门店店长发送巡检通知消息。",
     "1. 助手通过 dispatch query 或类似命令查询了已有派发任务\n2. 有可用任务则发送通知并反馈结果；无可用任务或缺少 form_url 则正确告知用户当前状态——以上均视为通过",
     "中", "case-07（需要有派发任务）", "葭牧"],
    # --- 报告类（高） ---
    ["case-11", "门店任务派发系统", "store-inspection", "门店巡检", "报告生成",
     "请使用 store-inspection 技能，帮我生成一份巡检汇总报告HTML文件。请先用 dispatch query 查看已有任务列表，选取最新任务，然后使用 report-generator 子技能的 generate-report 命令来生成多店汇总报告。",
     "1. 助手使用了 report-generator 子技能的 generate-report 命令\n2. 助手先通过 dispatch query 查询了可用任务并基于查询结果选取了正确的数据源\n3. 生成了HTML格式的报告文件，或如果没有可用任务数据则正确告知用户",
     "高", "case-07（需要有派发任务数据）", "葭牧"],
    ["case-12", "门店任务派发系统", "store-inspection", "门店巡检", "报告生成",
     "请使用 store-inspection 技能，帮我生成单店巡检报告HTML文件。请先用 store list 查看当前门店列表，然后选择第一家门店，使用 report-generator 子技能的 generate-report 命令以 single-store 模式生成报告。",
     "1. 助手尝试使用了 report-generator 子技能，或在无可用数据时正确告知用户\n2. 有数据则选择单店模式；无可用任务数据则正确告知用户也视为通过\n3. 生成了单店HTML报告；或无可用巡检数据时正确告知用户也视为通过",
     "高", "case-07（需要有派发任务数据）", "葭牧"],
    ["case-13", "门店任务派发系统", "store-inspection", "门店巡检", "报告推送",
     "请使用 store-inspection 技能，将位于 input/test_report.html 的巡检报告HTML文件上传到钉钉文档，并通过单条消息推送给我。我的 openDingTalkId 是 Du0y5VccnQiixQEldf67GrTZCOiSii7NBEbq，我叫葭牧（userId: 440535）。",
     "1. 助手使用 dws doc upload 上传 HTML 文件到钉钉文档\n2. 助手使用 dws doc permission add 为收件人添加了查看权限\n3. 助手发送了包含文档链接和文案摘要的单条消息",
     "高", "", "葭牧"],
    # --- 端到端/多轮（高） ---
    ["case-19", "门店任务派发系统", "store-inspection", "门店巡检", "端到端流程",
     "请使用 store-inspection 技能，帮我用系统中已有的巡检模板向所有门店派发一次新的巡检任务并通知到各门店店长，截止时间设为明天中午12:00。请一次性完成派发和通知。",
     "1. 助手通过 dispatch init 创建了派发任务实例\n2. 助手尝试通过 dispatch notify 发送通知，或在 form_url 因服务同步延迟无法获取时正确引导用户完成表单发布\n3. 助手在 extra_info 或命令参数中包含了截止时间信息\n4. 助手向用户反馈了当前操作进度和结果",
     "高", "case-05 + case-03（需要有模板和门店）", "葭牧"],
    ["case-20", "门店任务派发系统", "store-inspection", "门店巡检", "多轮交互",
     "请使用 store-inspection 技能，帮我创建一个新的巡检模板。",
     "1. 助手没有在信息不完整时直接执行创建操作，而是主动询问了必要信息（无论自然语言还是 ask_human 表单均可）\n2. 助手通过多轮对话逐步收集了模板名称和检查项等信息\n3. 助手在收集完信息后执行了模板创建或展示了完整配置方案等待确认（安全行为，也视为通过）",
     "中", "", "葭牧"],
]

body_font = Font(name="微软雅黑", size=10)
wrap_align = Alignment(vertical="top", wrap_text=True)

for r, row_data in enumerate(rows, 2):
    for c, val in enumerate(row_data, 1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.font = body_font
        cell.alignment = wrap_align
        cell.border = thin_border

col_widths = [12, 16, 22, 10, 10, 60, 80, 8, 30, 8]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

wb.save("evals/golden_cases_store_inspection.xlsx")
print("Excel generated successfully with 20 rows")
