#!/usr/bin/env python3
"""Generate a DOCX file from finalized labor contract text.

Usage:
  python3 scripts/generate_docx.py --input contract.md --output output/劳动合同.docx
  cat contract.md | python3 scripts/generate_docx.py --output output/劳动合同.docx

The script intentionally depends only on the finalized contract text. It does
not read or require the local development templates directory. When fed a
broader skill response, it keeps only the formal contract body and excludes
compliance notes, pending-field checklists, and DOCX action menus.
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt
except ImportError as exc:  # pragma: no cover - exercised only without dependency
    raise SystemExit(
        "Missing dependency: python-docx. Install it with: "
        "python3 -m pip install python-docx"
    ) from exc


TITLE_RE = re.compile(r"^劳动合同(?:（.*?）)?$")
MARKDOWN_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+)$")
CHAPTER_RE = re.compile(r"^(第[一二三四五六七八九十百]+[章节条]|[一二三四五六七八九十]+、).+")
ARTICLE_RE = re.compile(r"^第[一二三四五六七八九十百]+条\s*.+")
MARKDOWN_BOLD_RE = re.compile(r"\*\*(.*?)\*\*")
MARKDOWN_ITALIC_RE = re.compile(r"(?<!\*)\*(?!\*)(.*?)(?<!\*)\*(?!\*)")
MARKDOWN_CODE_RE = re.compile(r"`([^`]*)`")
MARKDOWN_LINK_RE = re.compile(r"\[([^\]]+)\]\([^)]+\)")
PLACEHOLDER_RE = re.compile(
    r"(人民币_{2,}元/月|第_{2,}种|_{2,}年_{2,}月_{2,}日|（\s*）|_{2,})"
)
KEY_CLAUSE_KEYWORDS = (
    "试用期",
    "劳动报酬",
    "保密",
    "知识产权",
    "竞业限制",
    "离职交接",
    "文书送达",
)
CONTRACT_START_RE = re.compile(r"^(?:#+\s*)?劳动合同(?:（.*?）)?$")
EXCLUDED_SECTION_MARKERS = (
    "【合规说明（简版）】",
    "待补充字段",
    "可选条款确认清单",
    "确认清单",
    "DOCX 选择",
    "编号式 DOCX 生成选择",
)
EMPTY_FIELD_RE = re.compile(r"(?P<label>[^\s：:]{1,30}[为是]：)\s*(?P<punct>[，。；])")
EMPTY_VALUE_AFTER_COLON_RE = re.compile(r"(?P<label>[^\s：:\n]{1,30}：)\s*(?P<punct>[，。；])")
LABEL_PREFIXES = (
    "甲方（用人单位）",
    "乙方（劳动者）",
    "名称",
    "统一社会信用代码",
    "注册地址",
    "法定代表人",
    "法定代表人或主要负责人",
    "联系电话",
    "姓名",
    "身份证号码",
    "身份证号码/护照号码/其他有效身份证件号码",
    "户籍地址",
    "有效通讯地址",
    "电子邮箱",
    "紧急联系人",
    "紧急联系人关系",
    "紧急联系人电话",
)


def set_run_font(run, size: int = 11, bold: bool = False, highlight=None) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    if highlight is not None:
        run.font.highlight_color = highlight
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def set_paragraph_spacing(paragraph, before: float = 0, after: float = 6, line: float = 1.25) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def set_document_defaults(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Pt(72)
    section.bottom_margin = Pt(72)
    section.left_margin = Pt(72)
    section.right_margin = Pt(72)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        style = styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.bold = True


def add_text_paragraph(document: Document, text: str, *, bold: bool = False, size: int = 11) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph)
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold)


def add_heading(document: Document, text: str, level: int) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, before=8, after=6, line=1.2)
    paragraph.style = document.styles[f"Heading {min(level, 3)}"]
    run = paragraph.add_run(text)
    highlight = WD_COLOR_INDEX.BRIGHT_GREEN if is_key_clause_heading(text) else None
    set_run_font(run, size=14 if level == 1 else 12, bold=True, highlight=highlight)


def add_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(paragraph, before=0, after=14, line=1.2)
    run = paragraph.add_run(text)
    set_run_font(run, size=18, bold=True)


def add_signature_line(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    set_paragraph_spacing(paragraph, before=2, after=4, line=1.2)
    add_segmented_runs(paragraph, text, default_size=11)


def clean_markdown_inline(text: str) -> str:
    """Remove common inline Markdown markers before writing text to Word."""
    cleaned = text.strip()
    previous = None
    while cleaned != previous:
        previous = cleaned
        cleaned = MARKDOWN_LINK_RE.sub(r"\1", cleaned)
        cleaned = MARKDOWN_CODE_RE.sub(r"\1", cleaned)
        cleaned = MARKDOWN_BOLD_RE.sub(r"\1", cleaned)
        cleaned = MARKDOWN_ITALIC_RE.sub(r"\1", cleaned)
    return cleaned


def is_key_clause_heading(text: str) -> bool:
    return any(keyword in text for keyword in KEY_CLAUSE_KEYWORDS)


def split_label_and_value(text: str) -> tuple[str, str] | None:
    if "：" not in text:
        return None
    label, value = text.split("：", 1)
    label = label.strip()
    value = value.lstrip()
    if not any(label.startswith(prefix) for prefix in LABEL_PREFIXES):
        return None
    return label + "：", value


def add_placeholder_aware_run(paragraph, text: str, *, bold: bool = False, size: int = 11) -> None:
    start = 0
    for match in PLACEHOLDER_RE.finditer(text):
        if match.start() > start:
            run = paragraph.add_run(text[start:match.start()])
            set_run_font(run, size=size, bold=bold)
        run = paragraph.add_run(match.group(0))
        set_run_font(run, size=size, bold=bold, highlight=WD_COLOR_INDEX.YELLOW)
        start = match.end()
    if start < len(text):
        run = paragraph.add_run(text[start:])
        set_run_font(run, size=size, bold=bold)


def add_segmented_runs(paragraph, text: str, *, default_size: int = 11) -> None:
    label_split = split_label_and_value(text)
    if label_split:
        label, value = label_split
        run = paragraph.add_run(label)
        set_run_font(run, size=default_size, bold=True)
        if value:
            add_placeholder_aware_run(paragraph, value, size=default_size)
        return
    add_placeholder_aware_run(paragraph, text, size=default_size)


def normalize_input(text: str) -> list[str]:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
    return [line.rstrip() for line in normalized.split("\n")]


def repair_empty_placeholders(text: str) -> str:
    """Patch common empty-field output into printable placeholders."""

    def replace_empty_field(match: re.Match[str]) -> str:
        return f"{match.group('label')}__________{match.group('punct')}"

    repaired = EMPTY_FIELD_RE.sub(replace_empty_field, text)
    repaired = EMPTY_VALUE_AFTER_COLON_RE.sub(replace_empty_field, repaired)
    return repaired


def extract_contract_text(text: str) -> str:
    """Extract only the formal contract body from a broader skill response.

    This keeps the finalized contract text while excluding strategy cards,
    compliance notes, pending-field checklists, and DOCX action menus.
    """
    lines = normalize_input(text)
    if not lines:
        return ""

    start_idx = 0
    for idx, raw_line in enumerate(lines):
        if CONTRACT_START_RE.match(raw_line.strip()):
            start_idx = idx
            break

    trimmed = lines[start_idx:]
    end_idx = len(trimmed)
    for idx, raw_line in enumerate(trimmed):
        line = clean_markdown_inline(raw_line.strip())
        if any(line.startswith(marker) for marker in EXCLUDED_SECTION_MARKERS):
            end_idx = idx
            break

    contract_lines = trimmed[:end_idx]
    while contract_lines and not contract_lines[-1].strip():
        contract_lines.pop()
    return repair_empty_placeholders("\n".join(contract_lines).strip())


def write_docx(text: str, output_path: Path) -> None:
    document = Document()
    set_document_defaults(document)

    lines = normalize_input(text)
    previous_blank = True

    for raw_line in lines:
        line = raw_line.strip()

        if not line:
            previous_blank = True
            continue

        markdown_heading = MARKDOWN_HEADING_RE.match(line)
        if markdown_heading:
            heading_level = len(markdown_heading.group(1))
            heading_text = clean_markdown_inline(markdown_heading.group(2))
            if TITLE_RE.match(heading_text):
                add_title(document, heading_text)
            else:
                add_heading(document, heading_text, min(heading_level, 3))
            previous_blank = False
            continue

        line = clean_markdown_inline(line)

        if TITLE_RE.match(line) and previous_blank:
            add_title(document, line)
        elif CHAPTER_RE.match(line) or ARTICLE_RE.match(line):
            add_heading(document, line, 2)
        elif line.startswith(("甲方（盖章）", "乙方（签字）", "法定代表人", "日期：")):
            add_signature_line(document, line)
        else:
            paragraph = document.add_paragraph()
            set_paragraph_spacing(paragraph)
            add_segmented_runs(paragraph, line)

        previous_blank = False

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate DOCX from finalized labor contract text.")
    parser.add_argument("--input", "-i", type=Path, help="Path to a UTF-8 text or Markdown contract file.")
    parser.add_argument("--output", "-o", type=Path, required=True, help="Output .docx path.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()

    if args.input:
        text = args.input.read_text(encoding="utf-8")
    else:
        text = sys.stdin.read()

    if not text.strip():
        raise SystemExit("No contract text provided. Use --input or pipe text through stdin.")

    contract_text = extract_contract_text(text)
    if not contract_text:
        raise SystemExit("No contract body detected after filtering. Provide finalized contract text.")

    write_docx(contract_text, args.output)
    print(f"DOCX generated: {args.output}")


if __name__ == "__main__":
    main()
