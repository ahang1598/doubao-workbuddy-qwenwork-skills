#!/usr/bin/env python3
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Pt, RGBColor

SCRIPT_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPT_DIR))
from validate_research import marker_path, sha256_file, validate_data  # noqa: E402


BODY_FONT = "宋体"
HEADING_FONT = "微软雅黑"
BLACK = "000000"
GRAY = "777777"
HEADER_FILL = "D0D0D0"
SUPPORT_FILL = "DCE6F1"
OTHER_FILL = "EEE8D5"
MISSING = "检索结果未载明"
NO_CASE_SENTENCE = "在本次检索及筛选范围内，未检索到符合条件的该类案例。"
SOURCE_NOTE = "注：本表仅列示经法智案例检索及案例全文阅读核验、且与本案争议焦点具有较高相似性的案例。案号可点击访问对应案例详情。"

CATEGORY_HEADERS = {
    "contract": "案涉合同名称",
    "project": "案涉项目名称",
    "financial": "案涉产品/计划名称",
    "transaction": "案涉交易名称",
    "other": "案涉标的",
}
TABLE_WIDTHS_CM = [1.0, 2.2, 2.7, 2.2, 2.7, 3.7, 10.1, 2.1]
DIRECTION_SECTIONS = [
    ("support", "（一）支持相关主张"),
    ("partial", "（二）部分支持或附条件支持相关主张"),
    ("oppose", "（三）不支持相关主张"),
]
PLATFORM_LABELS = {
    "case_search": "法智案例检索",
    "case_browser": "法智案例全文阅读",
    "legal_article_search": "法智法条综合检索",
    "law_content_visit": "法智法规全文阅读",
    "webpage_search": "法智联网搜索",
    "webpage_visit": "法智网页阅读",
}


def set_run_font(run, name=BODY_FONT, size=10.5, bold=None, color=BLACK):
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold


def format_paragraph(paragraph, before=0, after=6, line=1.35, align=None, first_line_cm=0.74):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.first_line_indent = Cm(first_line_cm) if first_line_cm else None
    if align is not None:
        paragraph.alignment = align


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def set_paragraph_bottom_border(paragraph, color="999999", size="4", space="5"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])
    set_run_font(run, BODY_FONT, 9, color=GRAY)


def add_hyperlink(paragraph, text, url, font=BODY_FONT, size=10.5, bold=False, color=BLACK):
    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia"):
        r_fonts.set(qn(f"w:{attr}"), font)
    r_pr.append(r_fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(round(size * 2)))
    r_pr.append(sz)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    r_pr.append(color_el)
    if bold:
        r_pr.append(OxmlElement("w:b"))
    run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35
    for level, size, before, after in ((1, 14, 14, 8), (2, 12, 10, 6), (3, 11, 8, 4)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = HEADING_FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_section(section, landscape=False):
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.top_margin = Cm(1.55)
        section.bottom_margin = Cm(1.45)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.35)
        section.right_margin = Cm(2.35)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.65)

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    clear_paragraph(hp)
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(4)
    run = hp.add_run("类案检索报告")
    set_run_font(run, HEADING_FONT, 10, bold=True)
    set_paragraph_bottom_border(hp)

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    clear_paragraph(fp)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(3)
    add_page_field(fp)


def clear_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def next_ids(numbering, tag, attr):
    values = [int(node.get(qn(attr))) for node in numbering.findall(qn(tag))]
    return max(values, default=-1) + 1


def new_decimal_num(doc, left=600, hanging=320):
    numbering = doc.part.numbering_part.element
    abstract_id = next_ids(numbering, "w:abstractNum", "w:abstractNumId")
    num_id = next_ids(numbering, "w:num", "w:numId")
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "decimal")
    level.append(num_fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level.append(level_text)
    suff = OxmlElement("w:suff")
    suff.set(qn("w:val"), "tab")
    level.append(suff)
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    p_pr.append(ind)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])
    p_pr.append(num_pr)


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.first_line_indent = None
    r = p.add_run(text)
    set_run_font(r, HEADING_FONT, {1: 14, 2: 12, 3: 11}[level], bold=True)
    return p


def add_body(doc, text, first_line=True, bold=False, after=6):
    p = doc.add_paragraph()
    format_paragraph(p, after=after, first_line_cm=0.74 if first_line else 0)
    r = p.add_run(text)
    set_run_font(r, BODY_FONT, 10.5, bold=bold)
    return p


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=68, bottom=68, start=85, end=85):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BLACK)
        borders.append(node)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_table_geometry(table):
    widths_dxa = [round(cm / 2.54 * 1440) for cm in TABLE_WIDTHS_CM]
    tbl_pr = table._tbl.tblPr
    width = tbl_pr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        tbl_pr.append(width)
    width.set(qn("w:w"), str(sum(widths_dxa)))
    width.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "0")
    indent.set(qn("w:type"), "dxa")
    tbl_pr.append(indent)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(value))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Cm(TABLE_WIDTHS_CM[index])
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")


def write_cell_text(cell, text, align, bold=False, size=8.5, font=BODY_FONT):
    cell.text = ""
    p = cell.paragraphs[0]
    format_paragraph(p, after=0, line=1.1, align=align, first_line_cm=0)
    r = p.add_run(text)
    set_run_font(r, font, size, bold=bold)


def write_cell_list(doc, cell, items):
    cell.text = ""
    num_id = new_decimal_num(doc, left=320, hanging=220)
    for index, text in enumerate(items):
        p = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        apply_numbering(p, num_id)
        format_paragraph(p, after=2 if index < len(items) - 1 else 0, line=1.08, align=WD_ALIGN_PARAGRAPH.LEFT, first_line_cm=0)
        r = p.add_run(text)
        set_run_font(r, BODY_FONT, 8.5)


def make_clean_doc(template_path=None):
    doc = Document(str(template_path)) if template_path and Path(template_path).exists() else Document()
    clear_body(doc)
    setup_styles(doc)
    configure_section(doc.sections[0], landscape=False)
    return doc


def build_report(data, template_path=None):
    doc = make_clean_doc(template_path)
    matter, search, cases = data["matter"], data["search"], data["cases"]

    # Cover.
    p = doc.add_paragraph()
    format_paragraph(p, before=36, after=0, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_cm=0)
    r = p.add_run(matter["title"])
    set_run_font(r, HEADING_FONT, 17, bold=True)
    if matter.get("case_number"):
        p = doc.add_paragraph()
        format_paragraph(p, before=18, after=0, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_cm=0)
        r = p.add_run(matter["case_number"])
        set_run_font(r, HEADING_FONT, 12, bold=True)
    p = doc.add_paragraph()
    format_paragraph(p, before=92, after=0, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_cm=0)
    r = p.add_run("类案检索报告")
    set_run_font(r, HEADING_FONT, 26, bold=True)
    for before, label, value in (
        (108, "检索日期", search["search_date"]),
        (28, "提交人", search["submitter"]),
        (12, "提交日期", search["submission_date"]),
    ):
        p = doc.add_paragraph()
        format_paragraph(p, before=before, after=0, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_cm=0)
        r = p.add_run(f"{label}：{value}")
        set_run_font(r, BODY_FONT, 11)
    doc.add_page_break()

    # Overview.
    add_heading(doc, "一、类案检索情况概述", 1)
    add_heading(doc, "（一）检索目的", 2)
    issues_text = "、".join(matter["issues"])
    add_body(doc, f"结合{matter['title']}（下称“本案”）的争议焦点，为查明司法实践对{issues_text}等问题的裁判观点，检索与本案基本事实和争议焦点具有较高相似性的案例，并对裁判观点进行归纳。")
    add_heading(doc, "（二）检索平台", 2)
    platform_text = "、".join(PLATFORM_LABELS[item] for item in search["platforms"])
    add_body(doc, platform_text + "。")
    add_heading(doc, "（三）检索条件", 2)
    search_num = new_decimal_num(doc)
    p = doc.add_paragraph()
    apply_numbering(p, search_num)
    format_paragraph(p, after=5, first_line_cm=0)
    r = p.add_run("检索关键词：")
    set_run_font(r, BODY_FONT, 10.5, bold=True)
    r = p.add_run("；".join(search["keywords"]) + "。")
    set_run_font(r)
    p = doc.add_paragraph()
    apply_numbering(p, search_num)
    format_paragraph(p, after=5, first_line_cm=0)
    r = p.add_run("筛选条件：")
    set_run_font(r, BODY_FONT, 10.5, bold=True)
    conditions = "；".join(search["screening_conditions"])
    r = p.add_run(f"时间范围为{search['time_scope']}，地域范围为{search['area_scope']}；{conditions}。")
    set_run_font(r)
    add_heading(doc, "（四）检索结果及结论", 2)
    counts = {key: sum(1 for case in cases if case["direction"] == key) for key, _ in DIRECTION_SECTIONS}
    result_text = (
        f"本次共初筛案例 {search['candidate_count']} 件，最终纳入与本案基本事实和争议焦点具有较高相似性的案例 {len(cases)} 件。"
        f"其中，支持相关主张 {counts['support']} 件，部分支持或附条件支持相关主张 {counts['partial']} 件，不支持相关主张 {counts['oppose']} 件。"
    )
    if len(cases) < 5:
        result_text += f"因{search['insufficient_reason']}，本报告按实际检索结果列示。"
    add_body(doc, result_text, bold=True)
    add_body(doc, search["result_conclusion"])
    p = doc.add_paragraph()
    format_paragraph(p, after=6, first_line_cm=0.74)
    r = p.add_run("本次检索核验的现行规范依据包括：")
    set_run_font(r)
    for index, rule in enumerate(data["legal_rules"]):
        if index:
            r = p.add_run("、")
            set_run_font(r)
        add_hyperlink(p, f"《{rule['law_name']}》{rule['article']}", rule["url"])
    r = p.add_run("。")
    set_run_font(r)

    # Comparison table section.
    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(landscape, landscape=True)
    add_heading(doc, "类案对照表", 1)
    headers = ["序号", "审理法院", "案件名称", "案号", CATEGORY_HEADERS[matter["category"]], "争议焦点", "法院观点", "裁判结论"]
    table = doc.add_table(rows=1, cols=8)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_borders(table)
    for index, label in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell)
        set_cell_shading(cell, HEADER_FILL)
        write_cell_text(cell, label, WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=9, font=HEADING_FONT)
    set_repeat_header(table.rows[0])
    for order, case in enumerate(cases, 1):
        cells = table.add_row().cells
        for cell in cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
        write_cell_text(cells[0], str(order), WD_ALIGN_PARAGRAPH.CENTER)
        write_cell_text(cells[1], case["court"], WD_ALIGN_PARAGRAPH.LEFT)
        write_cell_text(cells[2], case["title"], WD_ALIGN_PARAGRAPH.LEFT)
        cells[3].text = ""
        p = cells[3].paragraphs[0]
        format_paragraph(p, after=0, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_cm=0)
        add_hyperlink(p, case["case_no"], case["case_url"], size=8.5)
        write_cell_text(cells[4], case.get("object_name") or MISSING, WD_ALIGN_PARAGRAPH.LEFT)
        write_cell_list(doc, cells[5], case["issues"])
        write_cell_list(doc, cells[6], case["holdings"])
        set_cell_shading(cells[7], SUPPORT_FILL if case["outcome_class"] == "支持" else OTHER_FILL)
        cells[7].text = ""
        p = cells[7].paragraphs[0]
        format_paragraph(p, after=2, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_cm=0)
        r = p.add_run(case["outcome_class"])
        set_run_font(r, BODY_FONT, 8.5, bold=True)
        p = cells[7].add_paragraph()
        format_paragraph(p, after=0, line=1.1, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_cm=0)
        r = p.add_run(case["outcome_detail"])
        set_run_font(r, BODY_FONT, 8.5)
    set_table_geometry(table)
    p = doc.add_paragraph()
    format_paragraph(p, before=5, after=0, line=1.1, first_line_cm=0)
    r = p.add_run(SOURCE_NOTE)
    set_run_font(r, BODY_FONT, 8, color=GRAY)

    # Summary section.
    portrait = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(portrait, landscape=False)
    add_heading(doc, "二、类案裁判观点摘要", 1)
    summary_num = new_decimal_num(doc, left=420, hanging=300)
    for direction, heading in DIRECTION_SECTIONS:
        add_heading(doc, heading, 2)
        selected = [case for case in cases if case["direction"] == direction]
        if not selected:
            add_body(doc, NO_CASE_SENTENCE, first_line=False)
            continue
        for case in selected:
            p = doc.add_paragraph(style="Heading 3")
            apply_numbering(p, summary_num)
            p.paragraph_format.first_line_indent = None
            r = p.add_run(case["title"] + "（")
            set_run_font(r, HEADING_FONT, 11, bold=True)
            add_hyperlink(p, case["case_no"], case["case_url"], font=HEADING_FONT, size=11, bold=True)
            r = p.add_run("，" + case["court"] + "）")
            set_run_font(r, HEADING_FONT, 11, bold=True)
            add_body(doc, case["summary_fact"])
            add_body(doc, case["summary_reasoning"])

    add_heading(doc, "附：类案裁判文书目录", 1)
    appendix_num = new_decimal_num(doc)
    for case in cases:
        p = doc.add_paragraph()
        apply_numbering(p, appendix_num)
        format_paragraph(p, after=5, line=1.3, first_line_cm=0)
        r = p.add_run(case["title"] + "（")
        set_run_font(r)
        add_hyperlink(p, case["case_no"], case["case_url"])
        r = p.add_run("，" + case["court"] + "）。")
        set_run_font(r)

    doc.core_properties.title = matter["title"] + "类案检索报告"
    doc.core_properties.subject = "类案检索报告"
    doc.core_properties.author = search["submitter"]
    doc.core_properties.comments = "由 case-search-report 技能生成"
    return doc


def create_template(output_path: Path):
    placeholder = {
        "matter": {
            "title": "[案件名称]",
            "case_number": "[案号，可省略]",
            "category": "other",
            "issues": ["[争议焦点]"],
        },
        "search": {
            "search_date": "[YYYY-MM-DD]",
            "submitter": "[提交人]",
            "submission_date": "[YYYY-MM-DD]",
            "platforms": ["case_search", "case_browser", "legal_article_search", "law_content_visit"],
            "keywords": ["[关键词一]", "[关键词二]"],
            "date_from": "[YYYY-MM-DD]",
            "date_to": "[YYYY-MM-DD]",
            "time_scope": "[时间范围]",
            "area_scope": "[地域范围]",
            "screening_conditions": ["[筛选条件一]", "[筛选条件二]"],
            "candidate_count": 0,
            "insufficient_reason": "模板未填充研究数据",
            "result_conclusion": "[填写经真实案例全文和现行规范核验后形成的检索结论，不得加入律师建议或独立法律评价。]",
        },
        "legal_rules": [{"law_name": "规范名称", "article": "条款", "url": "https://www.fazhi.law"}],
        "cases": [],
    }
    doc = build_report(placeholder, template_path=None)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def audit_generated(path: Path, data):
    doc = Document(path)
    errors = []
    if len(doc.sections) != 3:
        errors.append("分节数量不是 3")
    else:
        expected = [WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE, WD_ORIENT.PORTRAIT]
        if [section.orientation for section in doc.sections] != expected:
            errors.append("分节方向不是纵向—横向—纵向")
    if len(doc.tables) != 1 or len(doc.tables[0].columns) != 8:
        errors.append("类案对照表不是唯一八列表格")
    else:
        header = [cell.text.strip() for cell in doc.tables[0].rows[0].cells]
        expected_header = ["序号", "审理法院", "案件名称", "案号", CATEGORY_HEADERS[data["matter"]["category"]], "争议焦点", "法院观点", "裁判结论"]
        if header != expected_header:
            errors.append("表头字段或顺序不一致")
        grid = doc.tables[0]._tbl.tblGrid.findall(qn("w:gridCol"))
        actual = [int(col.get(qn("w:w"))) for col in grid]
        expected = [round(cm / 2.54 * 1440) for cm in TABLE_WIDTHS_CM]
        if actual != expected:
            errors.append("表格列宽与固定规格不一致")
    texts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    required = [
        "一、类案检索情况概述", "（一）检索目的", "（二）检索平台", "（三）检索条件", "（四）检索结果及结论",
        "类案对照表", "二、类案裁判观点摘要", "（一）支持相关主张", "（二）部分支持或附条件支持相关主张",
        "（三）不支持相关主张", "附：类案裁判文书目录",
    ]
    positions = []
    for value in required:
        if value not in texts:
            errors.append(f"缺少固定标题：{value}")
        else:
            positions.append(texts.index(value))
    if positions and positions != sorted(positions):
        errors.append("固定标题顺序错误")
    prohibited = ("三、本案分析", "律师建议", "风险提示", "执行摘要", "法律依据")
    full_text = "\n".join(texts)
    for phrase in prohibited:
        if phrase in full_text:
            errors.append(f"出现禁止章节或文字：{phrase}")

    # Style, geometry and link audit. These checks intentionally inspect XML so
    # later template edits cannot silently loosen the fixed report contract.
    expected_style = {
        "Normal": (BODY_FONT, 21, False),
        "Heading 1": (HEADING_FONT, 28, True),
        "Heading 2": (HEADING_FONT, 24, True),
        "Heading 3": (HEADING_FONT, 22, True),
    }
    for style_name, (font_name, half_points, bold) in expected_style.items():
        style = doc.styles[style_name]
        r_pr = style._element.get_or_add_rPr()
        r_fonts = r_pr.find(qn("w:rFonts"))
        size = r_pr.find(qn("w:sz"))
        bold_node = r_pr.find(qn("w:b"))
        if r_fonts is None or r_fonts.get(qn("w:eastAsia")) != font_name:
            errors.append(f"{style_name} 中文字体错误")
        if size is None or int(size.get(qn("w:val"))) != half_points:
            errors.append(f"{style_name} 字号错误")
        if bool(bold_node is not None) != bold:
            errors.append(f"{style_name} 加粗设置错误")

    if len(doc.sections) == 3:
        for index, section in enumerate(doc.sections):
            if index == 1:
                expected_margins = (Cm(1.5), Cm(1.5))
                expected_size = (Cm(29.7), Cm(21.0))
            else:
                expected_margins = (Cm(2.35), Cm(2.35))
                expected_size = (Cm(21.0), Cm(29.7))
            if any(abs(int(actual) - int(expected)) > 1000 for actual, expected in zip((section.left_margin, section.right_margin), expected_margins)):
                errors.append(f"第 {index + 1} 节左右边距错误")
            if any(abs(int(actual) - int(expected)) > 1000 for actual, expected in zip((section.page_width, section.page_height), expected_size)):
                errors.append(f"第 {index + 1} 节页面尺寸错误")
            hp = section.header.paragraphs[0]
            if hp.alignment != WD_ALIGN_PARAGRAPH.RIGHT or hp.text.strip() != "类案检索报告":
                errors.append(f"第 {index + 1} 节页眉错误")
            if hp._p.find("./w:pPr/w:pBdr/w:bottom", hp._p.nsmap) is None:
                errors.append(f"第 {index + 1} 节页眉缺少底部细线")
            instr = "".join(node.text or "" for node in section.footer._element.findall(".//w:instrText", section.footer._element.nsmap))
            if "PAGE" not in instr:
                errors.append(f"第 {index + 1} 节页脚缺少页码域")

    if len(doc.tables) == 1 and len(doc.tables[0].columns) == 8:
        table = doc.tables[0]
        borders = table._tbl.find("./w:tblPr/w:tblBorders", table._tbl.nsmap)
        if borders is None:
            errors.append("表格缺少边框")
        else:
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                node = borders.find(qn(f"w:{edge}"))
                if node is None or node.get(qn("w:val")) != "single" or node.get(qn("w:sz")) != "4" or node.get(qn("w:color")) != BLACK:
                    errors.append(f"表格 {edge} 边框不符合黑色 0.5 磅单线")
        header_row = table.rows[0]
        if header_row._tr.find("./w:trPr/w:tblHeader", header_row._tr.nsmap) is None:
            errors.append("表头未设置跨页重复")
        if table._tbl.find("./w:tblPr/w:tblLayout", table._tbl.nsmap) is None:
            errors.append("表格未设置固定布局")
        if table._tbl.findall(".//w:trPr/w:trHeight", table._tbl.nsmap):
            errors.append("表格存在禁止的固定行高")
        for col_index, cell in enumerate(header_row.cells):
            shd = cell._tc.find("./w:tcPr/w:shd", cell._tc.nsmap)
            if shd is None or shd.get(qn("w:fill")) != HEADER_FILL:
                errors.append(f"表头第 {col_index + 1} 列填充错误")
            for run in cell._tc.findall(".//w:r", cell._tc.nsmap):
                if run.find(qn("w:t")) is None:
                    continue
                r_pr = run.find(qn("w:rPr"))
                fonts = r_pr.find(qn("w:rFonts")) if r_pr is not None else None
                size = r_pr.find(qn("w:sz")) if r_pr is not None else None
                if fonts is None or fonts.get(qn("w:eastAsia")) != HEADING_FONT:
                    errors.append(f"表头第 {col_index + 1} 列字体错误")
                    break
                if size is None or size.get(qn("w:val")) != "18":
                    errors.append(f"表头第 {col_index + 1} 列字号错误")
                    break
                if r_pr.find(qn("w:b")) is None:
                    errors.append(f"表头第 {col_index + 1} 列未加粗")
                    break
        for row_index, row in enumerate(table.rows[1:]):
            if row._tr.find("./w:trPr/w:trHeight", row._tr.nsmap) is not None:
                errors.append(f"案例行 {row_index + 1} 存在固定行高")
            for col_index, cell in enumerate(row.cells):
                margins = cell._tc.find("./w:tcPr/w:tcMar", cell._tc.nsmap)
                if margins is None:
                    errors.append(f"案例行 {row_index + 1} 第 {col_index + 1} 列缺少内边距")
                else:
                    expected_cell_margins = {"top": "68", "bottom": "68", "start": "85", "end": "85"}
                    for side, expected in expected_cell_margins.items():
                        node = margins.find(qn(f"w:{side}"))
                        if node is None or node.get(qn("w:w")) != expected or node.get(qn("w:type")) != "dxa":
                            errors.append(f"案例行 {row_index + 1} 第 {col_index + 1} 列 {side} 内边距错误")
                for run in cell._tc.findall(".//w:r", cell._tc.nsmap):
                    if run.find(qn("w:t")) is None:
                        continue
                    r_pr = run.find(qn("w:rPr"))
                    fonts = r_pr.find(qn("w:rFonts")) if r_pr is not None else None
                    size = r_pr.find(qn("w:sz")) if r_pr is not None else None
                    if fonts is None or fonts.get(qn("w:eastAsia")) != BODY_FONT:
                        errors.append(f"案例行 {row_index + 1} 第 {col_index + 1} 列字体错误")
                        break
                    if size is None or size.get(qn("w:val")) != "17":
                        errors.append(f"案例行 {row_index + 1} 第 {col_index + 1} 列字号错误")
                        break
            expected_fill = SUPPORT_FILL if data["cases"][row_index]["outcome_class"] == "支持" else OTHER_FILL
            shd = row.cells[7]._tc.find("./w:tcPr/w:shd", row.cells[7]._tc.nsmap)
            if shd is None or shd.get(qn("w:fill")) != expected_fill:
                errors.append(f"案例行 {row_index + 1} 裁判结论填充错误")

    if doc.element.body.findall(".//w:txbxContent", doc.element.body.nsmap) or doc.element.body.findall(".//w:pict", doc.element.body.nsmap) or doc.element.body.findall(".//w:drawing", doc.element.body.nsmap):
        errors.append("文档含禁止的文本框、艺术字或装饰图形")
    allowed_colors = {BLACK, GRAY}
    for color in doc.element.body.findall(".//w:color", doc.element.body.nsmap):
        value = color.get(qn("w:val"))
        if value and value not in allowed_colors:
            errors.append(f"正文出现未授权文字颜色：{value}")
            break
    allowed_fills = {HEADER_FILL, SUPPORT_FILL, OTHER_FILL}
    for shading in doc.element.body.findall(".//w:shd", doc.element.body.nsmap):
        value = shading.get(qn("w:fill"))
        if value and value not in {"auto", *allowed_fills}:
            errors.append(f"正文出现未授权填充颜色：{value}")
            break

    hyperlink_targets = []
    for hyperlink in doc.element.body.findall(".//w:hyperlink", doc.element.body.nsmap):
        rel_id = hyperlink.get(qn("r:id"))
        if rel_id and rel_id in doc.part.rels:
            hyperlink_targets.append(doc.part.rels[rel_id].target_ref)
    for case in data["cases"]:
        if hyperlink_targets.count(case["case_url"]) != 3:
            errors.append(f"案号链接未在表格、摘要和附录中各出现一次：{case['case_no']}")
        if full_text.count(case["title"]) < 2 or full_text.count(case["court"]) < 2:
            errors.append(f"案件信息在摘要和附录中不一致：{case['case_no']}")
    for rule in data["legal_rules"]:
        if hyperlink_targets.count(rule["url"]) != 1:
            errors.append(f"法条链接数量错误：{rule['law_name']}{rule['article']}")
    if errors:
        raise RuntimeError("DOCX 结构审计失败：" + "；".join(errors))


def main():
    parser = argparse.ArgumentParser(description="Build fixed-format case search report DOCX")
    parser.add_argument("--input", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--template", type=Path, default=SCRIPT_DIR.parent / "assets" / "case-search-report-template.docx")
    parser.add_argument("--create-template", type=Path)
    args = parser.parse_args()
    if args.create_template:
        create_template(args.create_template)
        print(args.create_template)
        return
    if not args.input or not args.output:
        parser.error("--input and --output are required unless --create-template is used")
    data = json.loads(args.input.read_text(encoding="utf-8"))
    errors = validate_data(data)
    if errors:
        raise RuntimeError("研究数据未通过校验：" + "；".join(errors))
    marker = marker_path(args.input)
    if not marker.exists():
        raise RuntimeError("缺少校验标记；先运行 validate_research.py")
    marker_data = json.loads(marker.read_text(encoding="utf-8"))
    if marker_data.get("sha256") != sha256_file(args.input):
        raise RuntimeError("研究 JSON 在校验后已被修改；请重新校验")
    if not args.template.exists():
        raise RuntimeError(f"模板不存在：{args.template}")
    doc = build_report(data, args.template)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    audit_generated(args.output, data)
    print(args.output)


if __name__ == "__main__":
    main()
