#!/usr/bin/env python3
"""Create the neutral, monochrome Word template used by this skill."""

from __future__ import annotations

import argparse
import zipfile
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT, WD_SECTION
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from lxml import etree
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Missing dependency. Install requirements.txt before running this script.") from exc

BLACK = RGBColor(0, 0, 0)
EAST_ASIA = {
    "Normal": "宋体",
    "Body Text": "宋体",
    "Title": "黑体",
    "Subtitle": "宋体",
    "Heading 1": "黑体",
    "Heading 2": "黑体",
    "Heading 3": "楷体",
    "Heading 4": "仿宋",
    "Cover Client": "黑体",
    "Cover Project": "黑体",
    "Cover Report": "黑体",
    "Cover Date": "宋体",
    "Letter Title": "黑体",
}

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W_NS}


def strip_cover_header_footer_refs(path: Path) -> None:
    """Remove first-section header/footer refs after python-docx serialization."""
    temporary = path.with_suffix(".rewrite.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "word/document.xml":
                root = etree.fromstring(data)
                sections = root.xpath(".//w:sectPr", namespaces=NS)
                if sections:
                    for reference in sections[0].xpath(
                        "./w:headerReference|./w:footerReference", namespaces=NS
                    ):
                        reference.getparent().remove(reference)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(info, data)
    temporary.replace(path)


def set_style_font(style, east_asia: str, size: float, bold: bool = False) -> None:
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = BLACK
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    for key, value in (("ascii", "Times New Roman"), ("hAnsi", "Times New Roman"), ("eastAsia", east_asia), ("cs", "Times New Roman")):
        rfonts.set(qn(f"w:{key}"), value)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")
    color = rpr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rpr.append(color)
    color.set(qn("w:val"), "000000")


def add_style(doc: Document, name: str, base: str, east_asia: str, size: float, bold: bool = False):
    styles = doc.styles
    if name in styles:
        style = styles[name]
    else:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = styles[base]
    set_style_font(style, east_asia, size, bold)
    return style


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    set_style_font(normal, EAST_ASIA["Normal"], 12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    body = doc.styles["Body Text"]
    set_style_font(body, EAST_ASIA["Body Text"], 12)
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.line_spacing = 1.5
    body.paragraph_format.first_line_indent = Pt(24)
    body.paragraph_format.space_before = Pt(0)
    body.paragraph_format.space_after = Pt(0)

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, EAST_ASIA["Heading 1"], 16, True)
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(9)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, EAST_ASIA["Heading 2"], 14, True)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = doc.styles["Heading 3"]
    set_style_font(h3, EAST_ASIA["Heading 3"], 12, True)
    h3.paragraph_format.space_before = Pt(9)
    h3.paragraph_format.space_after = Pt(3)
    h3.paragraph_format.keep_with_next = True

    h4 = doc.styles["Heading 4"]
    set_style_font(h4, EAST_ASIA["Heading 4"], 12, True)
    h4.paragraph_format.keep_with_next = True

    for style_name, east_asia, size in (
        ("Heading 1 Char", "黑体", 16),
        ("Heading 2 Char", "黑体", 14),
        ("Heading 3 Char", "楷体", 12),
        ("Heading 4 Char", "仿宋", 12),
    ):
        if style_name in doc.styles:
            set_style_font(doc.styles[style_name], east_asia, size, True)

    for name, base, font, size, bold, align in (
        ("Cover Client", "Normal", "黑体", 16, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("Cover Project", "Normal", "黑体", 16, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("Cover Report", "Normal", "黑体", 18, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("Cover Date", "Normal", "宋体", 12, False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Letter Title", "Normal", "黑体", 16, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("No Indent Body", "Body Text", "宋体", 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY),
        ("Signature", "Normal", "宋体", 12, False, WD_ALIGN_PARAGRAPH.LEFT),
    ):
        style = add_style(doc, name, base, font, size, bold)
        style.paragraph_format.alignment = align
        if name == "No Indent Body":
            style.paragraph_format.first_line_indent = Pt(0)
        if name.startswith("Cover"):
            style.paragraph_format.space_after = Pt(18)

    if "Legal Table" not in doc.styles:
        table_style = doc.styles.add_style("Legal Table", WD_STYLE_TYPE.TABLE)
    else:
        table_style = doc.styles["Legal Table"]
    set_style_font(table_style, "宋体", 10.5)


def configure_section(section, *, portrait: bool = True) -> None:
    section.orientation = WD_ORIENT.PORTRAIT if portrait else WD_ORIENT.LANDSCAPE
    section.page_width = Cm(21.0 if portrait else 29.7)
    section.page_height = Cm(29.7 if portrait else 21.0)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)


def set_page_number_start(section, value: int | None) -> None:
    sect_pr = section._sectPr
    node = sect_pr.find(qn("w:pgNumType"))
    if value is None:
        if node is not None:
            node.attrib.pop(qn("w:start"), None)
        return
    if node is None:
        node = OxmlElement("w:pgNumType")
        sect_pr.append(node)
    node.set(qn("w:start"), str(value))


def set_vertical_alignment(section, value: str) -> None:
    sect_pr = section._sectPr
    node = sect_pr.find(qn("w:vAlign"))
    if node is None:
        node = OxmlElement("w:vAlign")
        sect_pr.append(node)
    node.set(qn("w:val"), value)


def add_field(paragraph, instruction: str) -> None:
    paragraph.add_run()._r.append(OxmlElement("w:fldChar"))
    paragraph.runs[-1]._r[-1].set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    paragraph.add_run()._r.append(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    paragraph.add_run()._r.append(separate)
    paragraph.add_run("1")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)


def set_bottom_border(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pbdr.append(bottom)


def configure_header_footer(section, link_previous: bool, restart: int | None = None) -> None:
    section.header.is_linked_to_previous = link_previous
    section.footer.is_linked_to_previous = link_previous
    if not link_previous:
        hp = section.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.style = "No Indent Body"
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(15.0))
        hp.add_run("{{TARGET_SHORT_NAME}}")
        hp.add_run("\t")
        hp.add_run("{{REPORT_SHORT_TITLE}}")
        for run in hp.runs:
            run.font.size = Pt(10.5)
        set_bottom_border(hp)
        fp = section.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fp.style = "No Indent Body"
        add_field(fp, "PAGE")
    set_page_number_start(section, restart)


def add_marker(doc: Document, name: str) -> None:
    p = doc.add_paragraph(style="Body Text")
    p.add_run(f"{{{{BLOCK:{name}}}}}")


def build_template(output: Path) -> None:
    doc = Document()
    configure_styles(doc)
    cover = doc.sections[0]
    configure_section(cover)
    set_vertical_alignment(cover, "center")
    # Do not access cover.header/footer here. python-docx creates relationship
    # parts as soon as those properties are touched, which would make the
    # otherwise blank cover structurally reference a header/footer.

    doc.add_paragraph("{{CLIENT_NAME}}", style="Cover Client")
    doc.add_paragraph("关于{{TARGET_ENTITY}}的", style="Cover Project")
    doc.add_paragraph("法律尽职调查报告", style="Cover Report")
    doc.add_paragraph("{{REPORT_NO_LINE}}", style="Cover Date")
    doc.add_paragraph("{{REPORT_DATE}}", style="Cover Date")

    letter = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(letter)
    configure_header_footer(letter, False, 1)
    set_vertical_alignment(letter, "top")
    doc.add_paragraph("关于{{TARGET_ENTITY}}的法律尽职调查报告", style="Letter Title")
    doc.add_paragraph("致：{{CLIENT_NAME}}", style="Heading 2")
    add_marker(doc, "report_letter")
    doc.add_paragraph("免责声明", style="Heading 2")
    add_marker(doc, "disclaimer")

    body = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(body)
    configure_header_footer(body, True, None)
    set_vertical_alignment(body, "top")
    body_topics = (
        ("一、目标公司基本情况", "basic"),
        ("二、历史沿革", "history"),
        ("三、股权结构、股东及实际控制人", "ownership"),
        ("四、公司治理", "governance"),
        ("五、业务、资质及重大合同", "business"),
        ("六、资产与知识产权", "assets_ip"),
        ("七、劳动用工", "employment"),
        ("八、税务、诉讼、处罚及合规", "compliance"),
        ("九、其他实际启用模块", "other_modules"),
    )
    for title, marker in body_topics:
        doc.add_paragraph(title, style="Heading 1")
        add_marker(doc, marker)
    for title, marker in (
        ("重大问题及风险提示", "major_issues"),
        ("待核实事项", "pending_items"),
        ("结论与建议", "conclusion"),
        ("附件或资料清单", "attachments"),
    ):
        doc.add_paragraph(title, style="Heading 1")
        add_marker(doc, marker)

    signature = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(signature)
    configure_header_footer(signature, True, None)
    set_vertical_alignment(signature, "center")
    p = doc.add_paragraph("[本页为签署页，无正文]", style="Signature")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("出具主体：{{ISSUER}}", style="Signature")
    doc.add_paragraph("负责人：____________________", style="Signature")
    doc.add_paragraph("经办律师：__________________", style="Signature")
    doc.add_paragraph("经办律师：__________________", style="Signature")
    p = doc.add_paragraph("日期：{{SIGNATURE_DATE}}", style="Signature")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    core = doc.core_properties
    core.author = ""
    core.last_modified_by = ""
    for child in list(cover._sectPr):
        if child.tag in (qn("w:headerReference"), qn("w:footerReference")):
            cover._sectPr.remove(child)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    strip_cover_header_footer_refs(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build_template(args.output)


if __name__ == "__main__":
    main()
