#!/usr/bin/env python3
"""Fill the formal legal due-diligence DOCX template deterministically."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from pathlib import Path
from typing import Any

try:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from lxml import etree
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Missing dependency. Install requirements.txt before running this script.") from exc

BLACK = RGBColor(0, 0, 0)
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKGREL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
XML_NS = {"w": W_NS}
MARKER_RE = re.compile(r"^\{\{BLOCK:([a-z_]+)\}\}$")
FOOTNOTE_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
FOOTNOTE_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
BRAND_NAME = "同花顺旗下快查企业数据引擎"
BRAND_DISCLOSURE = (
    "本次公开信息核验以同花顺旗下快查企业数据引擎为主要渠道，并对官网、官方公示及可信公开报道"
    "进行有限补充核查；相关结果仅用于一致性复核和线索补充，不替代委托方提供并纳入当前采信范围的资料及原始证明。"
)
PUBLIC_PRE_DD_LIMITATION = (
    "本项目未提供可纳入当前采信资料库的项目资料，本报告属于公开信息 Pre-DD，"
    "仅在已核验公开信息的范围内形成初步判断，不替代全量法律尽职调查及原始证明核验。"
)
STATUS_WORDING = {
    "EMPTY_RESULT": "截至报告基准日，现有项目资料及已核验的公开信息未显示相关事项；该判断仍受资料完整性及信息更新时点限制。",
    "UNSUPPORTED": "现有公开信息不足以支持完整核验，相关事项暂不作无保留结论。",
    "FAILED": "本阶段尚未取得足以支持明确判断的资料，建议补充取得相关证明。",
    "INCOMPLETE": "现有公开信息样本不足以支持完整统计，相关数量及状态应以完整清单和原始证明为准。",
}
INTERNAL_LANGUAGE = {
    "MCP": re.compile(r"\bMCP\b", re.IGNORECASE),
    "discover": re.compile(r"\bdiscover\b", re.IGNORECASE),
    "call": re.compile(r"\bcall\b", re.IGNORECASE),
    "tool_id": re.compile(r"\btool_id\b", re.IGNORECASE),
    "查询日志": re.compile(r"查询日志"),
    "分页完整性": re.compile(r"分页完整性"),
    "参数错误": re.compile(r"参数错误"),
    "权限错误": re.compile(r"权限错误"),
    "服务错误": re.compile(r"服务错误"),
    "快查接口": re.compile(r"快查(?:查询|变更记录)?接口"),
    "快查能力": re.compile(r"快查能力"),
    "内部工作材料": re.compile(r"内部工作材料"),
    "快查": re.compile(r"快查"),
}
HEADING_FONTS = {
    "Heading1": "黑体",
    "Heading2": "黑体",
    "Heading3": "楷体",
    "Heading4": "仿宋",
}
BODY_KEYS = (
    "report_letter",
    "disclaimer",
    "basic",
    "history",
    "ownership",
    "governance",
    "business",
    "assets_ip",
    "employment",
    "compliance",
    "other_modules",
    "major_issues",
    "pending_items",
    "conclusion",
    "attachments",
)


def validate_source_control(data: dict[str, Any]) -> str:
    control = data.get("source_control")
    if not isinstance(control, dict):
        raise SystemExit("Missing required source_control object.")
    basis = control.get("primary_basis")
    if basis not in {"current_accepted_materials", "public_information_pre_dd"}:
        raise SystemExit("source_control.primary_basis must be current_accepted_materials or public_information_pre_dd.")
    material_ids = control.get("primary_material_ids")
    if not isinstance(material_ids, list):
        raise SystemExit("source_control.primary_material_ids must be an array.")
    if basis == "current_accepted_materials" and not any(str(value).strip() for value in material_ids):
        raise SystemExit("current_accepted_materials requires non-empty primary_material_ids.")
    if basis == "public_information_pre_dd" and material_ids:
        raise SystemExit("public_information_pre_dd requires empty primary_material_ids.")
    if control.get("external_verification_role") != "corroboration_only":
        raise SystemExit("source_control.external_verification_role must be corroboration_only.")
    if control.get("brand_disclosure_required") is not True:
        raise SystemExit("source_control.brand_disclosure_required must be true.")
    if not isinstance(control.get("conflicts"), list):
        raise SystemExit("source_control.conflicts must be an array.")
    return basis


def clientize_text(value: str) -> str:
    result = value
    for status, wording in STATUS_WORDING.items():
        result = re.sub(rf"\b{status}\b", wording, result)
    return result


def clientize_value(value: Any) -> Any:
    if isinstance(value, str):
        return clientize_text(value)
    if isinstance(value, list):
        return [clientize_value(item) for item in value]
    if isinstance(value, dict):
        return {key: clientize_value(item) for key, item in value.items()}
    return value


def visible_strings(value: Any, location: str):
    if isinstance(value, str):
        yield location, value
    elif isinstance(value, list):
        for index, item in enumerate(value):
            yield from visible_strings(item, f"{location}[{index}]")
    elif isinstance(value, dict):
        for key, item in value.items():
            yield from visible_strings(item, f"{location}.{key}")


def reject_internal_language(data: dict[str, Any]) -> None:
    visible = {key: data.get(key) for key in BODY_KEYS}
    visible["meta"] = data.get("meta", {})
    visible["signature"] = data.get("signature", {})
    for location, text in visible_strings(visible, "report"):
        for label, pattern in INTERNAL_LANGUAGE.items():
            if pattern.search(text):
                raise SystemExit(f"Internal client-facing language detected at {location}: {label}")


def append_report_letter_statement(value: Any, statement: str) -> list[Any]:
    if value is None or value == "":
        return [statement]
    if isinstance(value, list):
        return [*value, statement]
    return [value, statement]


def set_explicit_rfonts(rpr, east_asia: str) -> None:
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    for attr, value in (
        ("ascii", "Times New Roman"),
        ("hAnsi", "Times New Roman"),
        ("eastAsia", east_asia),
        ("cs", "Times New Roman"),
    ):
        rfonts.set(qn(f"w:{attr}"), value)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")


def normalize_heading_fonts(doc: Document) -> None:
    for style_id, east_asia in HEADING_FONTS.items():
        for candidate in (style_id, f"{style_id}Char"):
            matches = doc.styles.element.xpath(f"./w:style[@w:styleId='{candidate}']")
            if not matches:
                continue
            style = matches[0]
            rpr = style.find(qn("w:rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                style.append(rpr)
            set_explicit_rfonts(rpr, east_asia)
    for paragraph in doc.paragraphs:
        style_id = paragraph.style.style_id if paragraph.style is not None else ""
        east_asia = HEADING_FONTS.get(style_id)
        if east_asia is None:
            continue
        for run in paragraph.runs:
            set_explicit_rfonts(run._r.get_or_add_rPr(), east_asia)


def strip_cover_header_footer_refs(path: Path) -> None:
    """Keep the cover structurally free of header and footer references."""
    temporary = path.with_suffix(".rewrite.docx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        temporary, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "word/document.xml":
                root = etree.fromstring(data)
                sections = root.xpath(".//w:sectPr", namespaces=XML_NS)
                if sections:
                    for reference in sections[0].xpath(
                        "./w:headerReference|./w:footerReference", namespaces=XML_NS
                    ):
                        reference.getparent().remove(reference)
                data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(info, data)
    temporary.replace(path)


def set_run_font(run, east_asia: str = "宋体", size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run.font.color.rgb = BLACK
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), east_asia)
    rfonts.set(qn("w:cs"), "Times New Roman")
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")


def set_run_black(run) -> None:
    run.font.color.rgb = BLACK
    rpr = run._r.get_or_add_rPr()
    color = rpr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rpr.append(color)
    color.set(qn("w:val"), "000000")


def replace_paragraph_text(paragraph, replacements: dict[str, str]) -> None:
    original = "".join(run.text for run in paragraph.runs) or paragraph.text
    updated = original
    for key, value in replacements.items():
        updated = updated.replace("{{" + key + "}}", value)
    if updated == original:
        return
    for run in paragraph.runs:
        run.text = ""
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    run.text = updated
    set_run_black(run)


def all_paragraphs(doc: Document):
    yield from doc.paragraphs
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def move_before(element, marker_paragraph) -> None:
    marker_paragraph._p.addprevious(element)


def format_external_source(source: Any) -> str:
    if not isinstance(source, dict):
        raise SystemExit("external_sources 每项必须为对象。")
    required = ("source_name", "title", "accessed_on", "url")
    missing = [key for key in required if not str(source.get(key, "")).strip()]
    if missing:
        raise SystemExit("external_sources 缺少字段：" + ", ".join(missing))
    url = str(source["url"]).strip()
    if not re.match(r"^https?://", url, re.IGNORECASE):
        raise SystemExit("external_sources.url 必须为 http 或 https 地址。")
    published = str(source.get("published_on") or "未载明").strip()
    accessed = str(source["accessed_on"]).strip()
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", accessed):
        raise SystemExit("external_sources.accessed_on 必须使用 YYYY-MM-DD。")
    if published != "未载明" and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", published):
        raise SystemExit("external_sources.published_on 必须使用 YYYY-MM-DD 或留空。")
    return (
        f"来源：{str(source['source_name']).strip()}，《{str(source['title']).strip()}》，"
        f"发布日期：{published}，{url}（访问日期：{accessed}）。"
    )


def add_paragraph_before(
    doc: Document,
    marker,
    text: str,
    style: str = "Body Text",
    external_sources: Any = None,
    footnotes: list[tuple[str, str]] | None = None,
):
    paragraph = doc.add_paragraph(style=style)
    paragraph.add_run(text)
    if external_sources not in (None, []):
        if style not in ("Body Text", "No Indent Body"):
            raise SystemExit("external_sources 只能附加到正文段落。")
        if not isinstance(external_sources, list):
            raise SystemExit("external_sources 必须为数组。")
        if footnotes is None:
            raise RuntimeError("Footnote collector is required for external sources.")
        for source in external_sources:
            marker_text = f"[[LEGAL_DD_FN_{len(footnotes) + 1:04d}]]"
            paragraph.add_run(marker_text)
            footnotes.append((marker_text, format_external_source(source)))
    move_before(paragraph._p, marker)
    return paragraph


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), "000000")


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    node = tr_pr.find(qn("w:tblHeader"))
    if node is None:
        node = OxmlElement("w:tblHeader")
        tr_pr.append(node)
    node.set(qn("w:val"), "true")


def set_cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = tc_pr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tc_pr.append(mar)
    for edge, value in (("top", 85), ("bottom", 85), ("left", 113), ("right", 113)):
        node = mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_cm: list[float]) -> None:
    widths = [int(round(value * 567.0)) for value in widths_cm]
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            cell.width = Cm(widths_cm[min(idx, len(widths_cm) - 1)])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_table_before(doc: Document, marker, headers: list[str], rows: list[list[Any]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Legal Table"
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = str(text)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_repeat_header(table.rows[0])
    for row_data in rows:
        row = table.add_row()
        for idx in range(len(headers)):
            row.cells[idx].text = str(row_data[idx]) if idx < len(row_data) and row_data[idx] is not None else ""
            row.cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths is None:
        widths = [15.0 / len(headers)] * len(headers)
    set_table_geometry(table, widths)
    set_table_borders(table)
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                paragraph.style = "No Indent Body"
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.15
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_idx == 0 or col_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_run_font(run, "黑体" if row_idx == 0 else "宋体", 10.5, row_idx == 0)
    move_before(table._tbl, marker)
    return table


def normalize_blocks(key: str, value: Any) -> list[dict[str, Any]]:
    if value is None or value == [] or value == "":
        return [{"type": "paragraph", "text": "本次调查未取得足以支持进一步结论的资料，相关事项列入待核实范围。"}]
    if isinstance(value, str):
        return [{"type": "paragraph", "text": value}]
    if isinstance(value, dict):
        return [value]
    if not isinstance(value, list):
        return [{"type": "paragraph", "text": str(value)}]
    if key == "major_issues" and value and isinstance(value[0], dict) and "facts" in value[0]:
        blocks: list[dict[str, Any]] = []
        for idx, item in enumerate(value, 1):
            number = item.get("number") or f"DD-{idx:02d}"
            blocks.extend([
                {"type": "heading2", "text": f"（{idx}）{number}  {item.get('title', '重大问题')}"},
                {
                    "type": "paragraph",
                    "text": f"事实与证据状态：{item.get('facts', '')}；{item.get('evidence_status', '待核实')}。",
                    "external_sources": item.get("external_sources", []),
                },
                {"type": "paragraph", "text": f"风险及影响：{item.get('risk', '')}。{item.get('impact', '')}"},
                {"type": "paragraph", "text": f"建议处理路径：{item.get('recommendation', '')}"},
            ])
        return blocks
    if key == "pending_items" and value and isinstance(value[0], dict):
        rows = [[x.get("priority", ""), x.get("item", ""), x.get("reason", ""), x.get("related_issue", "")] for x in value]
        return [{"type": "table", "headers": ["优先级", "待核实事项", "原因或所需材料", "关联问题"], "rows": rows, "widths_cm": [2.0, 5.0, 5.5, 2.5]}]
    if key == "attachments" and value and isinstance(value[0], dict):
        rows = [[x.get("number", i), x.get("name", ""), x.get("source", ""), x.get("status", "")] for i, x in enumerate(value, 1)]
        return [{"type": "table", "headers": ["序号", "资料或附件名称", "来源", "状态"], "rows": rows, "widths_cm": [1.5, 7.5, 3.5, 2.5]}]
    blocks = []
    for item in value:
        blocks.append({"type": "paragraph", "text": item} if isinstance(item, str) else item)
    return blocks


def fill_marker(doc: Document, marker, key: str, value: Any, footnotes: list[tuple[str, str]]) -> None:
    for block in normalize_blocks(key, value):
        block_type = block.get("type", "paragraph")
        if block_type == "table":
            add_table_before(doc, marker, list(block.get("headers", [])), list(block.get("rows", [])), block.get("widths_cm"))
            continue
        style = {"heading2": "Heading 2", "heading3": "Heading 3", "heading4": "Heading 4", "paragraph": "Body Text", "no_indent": "No Indent Body"}.get(block_type, "Body Text")
        paragraph = add_paragraph_before(
            doc,
            marker,
            str(block.get("text", "")),
            style,
            block.get("external_sources"),
            footnotes,
        )
        if block.get("keep_with_next"):
            paragraph.paragraph_format.keep_with_next = True
    marker._p.getparent().remove(marker._p)


def scrub_properties(doc: Document) -> None:
    core = doc.core_properties
    core.author = ""
    core.last_modified_by = ""
    core.comments = ""
    core.keywords = ""
    core.subject = ""


def xml_bytes(root) -> bytes:
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")


def next_relationship_id(rels_root) -> str:
    maximum = 0
    for rel in rels_root.findall(f"{{{PKGREL_NS}}}Relationship"):
        match = re.fullmatch(r"rId(\d+)", rel.get("Id") or "")
        if match:
            maximum = max(maximum, int(match.group(1)))
    return f"rId{maximum + 1}"


def ensure_footnote_relationship(rels_root) -> None:
    for rel in rels_root.findall(f"{{{PKGREL_NS}}}Relationship"):
        if rel.get("Type") == FOOTNOTE_REL_TYPE:
            return
    rel = etree.SubElement(rels_root, f"{{{PKGREL_NS}}}Relationship")
    rel.set("Id", next_relationship_id(rels_root))
    rel.set("Type", FOOTNOTE_REL_TYPE)
    rel.set("Target", "footnotes.xml")


def ensure_footnote_content_type(content_types_root) -> None:
    part_name = "/word/footnotes.xml"
    for override in content_types_root.findall(f"{{{CT_NS}}}Override"):
        if override.get("PartName") == part_name:
            override.set("ContentType", FOOTNOTE_CONTENT_TYPE)
            return
    override = etree.SubElement(content_types_root, f"{{{CT_NS}}}Override")
    override.set("PartName", part_name)
    override.set("ContentType", FOOTNOTE_CONTENT_TYPE)


def make_footnotes_root():
    root = etree.Element(f"{{{W_NS}}}footnotes", nsmap={"w": W_NS, "r": R_NS})
    for note_id, note_type, separator in (
        ("-1", "separator", "separator"),
        ("0", "continuationSeparator", "continuationSeparator"),
    ):
        note = etree.SubElement(root, f"{{{W_NS}}}footnote")
        note.set(f"{{{W_NS}}}id", note_id)
        note.set(f"{{{W_NS}}}type", note_type)
        paragraph = etree.SubElement(note, f"{{{W_NS}}}p")
        run = etree.SubElement(paragraph, f"{{{W_NS}}}r")
        etree.SubElement(run, f"{{{W_NS}}}{separator}")
    return root


def next_footnote_id(footnotes_root) -> int:
    ids: list[int] = []
    for note in footnotes_root.findall(f"{{{W_NS}}}footnote"):
        try:
            value = int(note.get(f"{{{W_NS}}}id", ""))
        except ValueError:
            continue
        if value >= 1:
            ids.append(value)
    return max(ids, default=0) + 1


def add_black_note_run(paragraph, text: str, *, reference: bool = False):
    run = etree.SubElement(paragraph, f"{{{W_NS}}}r")
    properties = etree.SubElement(run, f"{{{W_NS}}}rPr")
    fonts = etree.SubElement(properties, f"{{{W_NS}}}rFonts")
    fonts.set(f"{{{W_NS}}}ascii", "Times New Roman")
    fonts.set(f"{{{W_NS}}}hAnsi", "Times New Roman")
    fonts.set(f"{{{W_NS}}}eastAsia", "宋体")
    color = etree.SubElement(properties, f"{{{W_NS}}}color")
    color.set(f"{{{W_NS}}}val", "000000")
    size = etree.SubElement(properties, f"{{{W_NS}}}sz")
    size.set(f"{{{W_NS}}}val", "18")
    size_cs = etree.SubElement(properties, f"{{{W_NS}}}szCs")
    size_cs.set(f"{{{W_NS}}}val", "18")
    if reference:
        vertical = etree.SubElement(properties, f"{{{W_NS}}}vertAlign")
        vertical.set(f"{{{W_NS}}}val", "superscript")
        etree.SubElement(run, f"{{{W_NS}}}footnoteRef")
    else:
        node = etree.SubElement(run, f"{{{W_NS}}}t")
        node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        node.text = text
    return run


def append_footnote(footnotes_root, note_id: int, text: str) -> None:
    note = etree.SubElement(footnotes_root, f"{{{W_NS}}}footnote")
    note.set(f"{{{W_NS}}}id", str(note_id))
    paragraph = etree.SubElement(note, f"{{{W_NS}}}p")
    properties = etree.SubElement(paragraph, f"{{{W_NS}}}pPr")
    spacing = etree.SubElement(properties, f"{{{W_NS}}}spacing")
    spacing.set(f"{{{W_NS}}}after", "0")
    spacing.set(f"{{{W_NS}}}line", "240")
    spacing.set(f"{{{W_NS}}}lineRule", "auto")
    add_black_note_run(paragraph, "", reference=True)
    add_black_note_run(paragraph, " " + text)


def insert_footnote_reference(document_root, marker: str, note_id: int) -> None:
    for text_node in document_root.xpath(".//w:t", namespaces=XML_NS):
        if text_node.text != marker:
            continue
        run = text_node.getparent()
        while run is not None and run.tag != f"{{{W_NS}}}r":
            run = run.getparent()
        if run is None:
            break
        parent = run.getparent()
        index = parent.index(run)
        parent.remove(run)
        reference_run = etree.Element(f"{{{W_NS}}}r")
        properties = etree.SubElement(reference_run, f"{{{W_NS}}}rPr")
        color = etree.SubElement(properties, f"{{{W_NS}}}color")
        color.set(f"{{{W_NS}}}val", "000000")
        vertical = etree.SubElement(properties, f"{{{W_NS}}}vertAlign")
        vertical.set(f"{{{W_NS}}}val", "superscript")
        reference = etree.SubElement(reference_run, f"{{{W_NS}}}footnoteReference")
        reference.set(f"{{{W_NS}}}id", str(note_id))
        parent.insert(index, reference_run)
        return
    raise SystemExit(f"Footnote marker missing from generated DOCX: {marker}")


def inject_footnotes(path: Path, footnotes: list[tuple[str, str]]) -> None:
    if not footnotes:
        return
    temporary = path.with_suffix(".footnotes.tmp.docx")
    with zipfile.ZipFile(path, "r") as source:
        names = {info.filename for info in source.infolist()}
        document_root = etree.fromstring(source.read("word/document.xml"))
        relationships_root = etree.fromstring(source.read("word/_rels/document.xml.rels"))
        content_types_root = etree.fromstring(source.read("[Content_Types].xml"))
        footnotes_root = (
            etree.fromstring(source.read("word/footnotes.xml"))
            if "word/footnotes.xml" in names
            else make_footnotes_root()
        )
        for marker, text in footnotes:
            note_id = next_footnote_id(footnotes_root)
            insert_footnote_reference(document_root, marker, note_id)
            append_footnote(footnotes_root, note_id, text)
        ensure_footnote_relationship(relationships_root)
        ensure_footnote_content_type(content_types_root)
        overrides = {
            "word/document.xml": xml_bytes(document_root),
            "word/_rels/document.xml.rels": xml_bytes(relationships_root),
            "[Content_Types].xml": xml_bytes(content_types_root),
            "word/footnotes.xml": xml_bytes(footnotes_root),
        }
        with zipfile.ZipFile(temporary, "w", compression=zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                target.writestr(info, overrides.get(info.filename, source.read(info.filename)))
            for name, content in overrides.items():
                if name not in names:
                    target.writestr(name, content)
    temporary.replace(path)


def build(template: Path, data_path: Path, output: Path) -> None:
    if not template.exists() or template.suffix.lower() != ".docx":
        raise SystemExit(f"Template not found or not DOCX: {template}")
    if output.exists():
        raise SystemExit(f"Refusing to overwrite existing report: {output}")
    data = json.loads(data_path.read_text(encoding="utf-8"))
    basis = validate_source_control(data)
    for key in BODY_KEYS:
        data[key] = clientize_value(data.get(key))
    data["meta"] = clientize_value(data.get("meta", {}))
    data["signature"] = clientize_value(data.get("signature", {}))
    reject_internal_language(data)
    if basis == "public_information_pre_dd":
        data["report_letter"] = append_report_letter_statement(data.get("report_letter"), PUBLIC_PRE_DD_LIMITATION)
    data["report_letter"] = append_report_letter_statement(data.get("report_letter"), BRAND_DISCLOSURE)
    meta = data.get("meta", {})
    required = ("client_name", "target_entity", "report_date", "cutoff_date")
    missing = [key for key in required if not str(meta.get(key, "")).strip()]
    if missing:
        raise SystemExit(f"Missing required meta fields: {', '.join(missing)}")
    target = str(meta["target_entity"])
    replacements = {
        "CLIENT_NAME": str(meta["client_name"]),
        "TARGET_ENTITY": target,
        "TARGET_SHORT_NAME": str(meta.get("target_short_name") or target),
        "REPORT_SHORT_TITLE": str(meta.get("report_short_title") or "法律尽职调查报告"),
        "REPORT_NO_LINE": ("报告文号：" + str(meta["report_no"])) if meta.get("report_no") else "",
        "REPORT_DATE": str(meta["report_date"]),
        "ISSUER": str(data.get("signature", {}).get("issuer") or "____________________"),
        "SIGNATURE_DATE": str(data.get("signature", {}).get("date") or meta["report_date"]),
    }
    doc = Document(template)
    footnotes: list[tuple[str, str]] = []
    normalize_heading_fonts(doc)
    for paragraph in all_paragraphs(doc):
        replace_paragraph_text(paragraph, replacements)
    markers = []
    for paragraph in list(doc.paragraphs):
        match = MARKER_RE.match(paragraph.text.strip())
        if match:
            markers.append((paragraph, match.group(1)))
    seen = {key for _, key in markers}
    missing_markers = [key for key in BODY_KEYS if key not in seen]
    if missing_markers:
        raise SystemExit(f"Template is missing block markers: {', '.join(missing_markers)}")
    for marker, key in markers:
        fill_marker(doc, marker, key, data.get(key), footnotes)
    normalize_heading_fonts(doc)
    scrub_properties(doc)
    for paragraph in all_paragraphs(doc):
        for run in paragraph.runs:
            set_run_black(run)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    strip_cover_header_footer_refs(output)
    inject_footnotes(output, footnotes)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", required=True, type=Path)
    parser.add_argument("--data", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    build(args.template, args.data, args.output)


if __name__ == "__main__":
    main()
