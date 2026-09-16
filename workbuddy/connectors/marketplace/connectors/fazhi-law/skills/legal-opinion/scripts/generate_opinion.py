# -*- coding: utf-8 -*-
"""
Generate the legal opinion DOCX for the factoring financing matter.
Based on the general special opinion template.
"""

import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml
import copy
import os
import argparse
from pathlib import Path

SKILL_ROOT = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = SKILL_ROOT / 'assets' / 'templates' / '01_通用专项法律意见书模板.docx'
DEFAULT_OUTPUT_NAME = '保理融资专项法律意见书.docx'

# ── helpers ──────────────────────────────────────────────

def set_cell_font(cell, font_name='楷体', font_size=Pt(12), bold=False):
    for p in cell.paragraphs:
        for run in p.runs:
            run.font.name = font_name
            run.font.size = font_size
            run.font.bold = bold
            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.find(qn('w:rFonts'))
            if rFonts is None:
                rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
                rPr.insert(0, rFonts)
            rFonts.set(qn('w:eastAsia'), font_name)

def add_paragraph_to_body(doc, text, style='Normal', bold=False, alignment=None, first_line_indent=None):
    """Add a paragraph with the given text and style."""
    p = doc.add_paragraph(text, style=style)
    if bold:
        for run in p.runs:
            run.font.bold = True
    if alignment is not None:
        p.alignment = alignment
    if first_line_indent is not None:
        p.paragraph_format.first_line_indent = first_line_indent
    return p

def set_run_font(run, name='楷体', size=Pt(12), bold=False, color=None):
    if size is not None:
        run.font.size = size
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), name)
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    if name:
        run.font.name = name
    if color:
        run.font.color.rgb = color

def add_body_paragraph(doc, text, indent=True, bold=False):
    """Add a body text paragraph with proper formatting."""
    p = doc.add_paragraph()
    p.style = doc.styles['Normal']
    run = p.add_run(text)
    set_run_font(run, '楷体', Pt(12), bold)
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.85)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    return p

def add_law_provision(doc, text):
    """Add a law provision paragraph."""
    p = doc.add_paragraph()
    p.style = doc.styles['Normal']
    run = p.add_run(text)
    set_run_font(run, '楷体', Pt(12))
    p.paragraph_format.first_line_indent = Cm(0.85)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(3)
    return p

def add_heading(doc, text, level=1):
    """Add a heading with the template's heading style."""
    p = doc.add_paragraph(text, style=f'Heading {level}')
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p

def add_empty_paragraph(doc):
    p = doc.add_paragraph()
    return p

# ── main generation ─────────────────────────────────────

def generate(output_path):
    # Open template to inherit styles
    doc = Document(str(TEMPLATE_PATH))

    # Clear all existing body content (paragraphs and tables)
    body = doc.element.body
    for child in list(body):
        if child.tag.endswith('}p') or child.tag.endswith('}tbl'):
            body.remove(child)
    # Keep sectPr (page settings)

    # ═══════════════════════════════════════════════════
    # COVER PAGE
    # ═══════════════════════════════════════════════════
    for _ in range(4):
        add_empty_paragraph(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('关于上海景和商业保理有限公司拟向\n浙江宏达供应链有限公司提供保理融资\n之专项法律意见书')
    set_run_font(run, '宋体', Pt(22), bold=True)
    p.paragraph_format.space_after = Pt(24)

    for _ in range(3):
        add_empty_paragraph(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('【待补充：律所或法务部门名称】')
    set_run_font(run, '宋体', Pt(14))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('二〇二六年八月十七日')
    set_run_font(run, '宋体', Pt(14))

    doc.add_page_break()

    # ═══════════════════════════════════════════════════
    # IMPORTANT DECLARATIONS
    # ═══════════════════════════════════════════════════
    add_heading(doc, '重要声明', 1)

    declarations = [
        '本声明为本法律意见书不可分割的组成部分。请在使用、转发或据此决策前完整阅读。',
        'AI辅助生成：本法律意见书由人工智能工具辅助形成。人工智能可能发生事实遗漏、规则误引、时效判断偏差或表达歧义，不能替代具备相应专业能力人员的独立判断。',
        '材料依赖：本法律意见书仅以用户提供的文件、数据、陈述和公开检索结果为基础。未经独立核验的材料不视为已经本所或出具主体确认。',
        '复核要求：凡涉及重大交易、监管申报、行政许可、争议处置、重大个人信息处理或可能产生重大损失的事项，均应由执业律师或有权专业人员复核后使用。',
        '身份与签章：非律师用户不得冒用律师事务所、律师姓名、执业证号或印章。未完成签发程序的文本不得对外宣称为律师事务所出具的正式法律意见。',
        '用途限制：本法律意见书仅供载明的委托事项和接收对象使用。任何第三方不得仅凭本文件替代自身调查、判断或取得专业意见。',
    ]
    for text in declarations:
        add_body_paragraph(doc, text, indent=True)

    doc.add_page_break()

    # ═══════════════════════════════════════════════════
    # TABLE OF CONTENTS (real TOC field)
    # ═══════════════════════════════════════════════════
    add_heading(doc, '目录', 1)
    p = doc.add_paragraph()
    run = p.add_run()
    set_run_font(run, '楷体', Pt(12))
    fldChar1 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="begin" w:dirty="true"/>')
    instrText = parse_xml(f'<w:instrText {nsdecls("w")} xml:space="preserve"> TOC \\o "1-2" \\h \\z \\u </w:instrText>')
    fldChar2 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="separate"/>')
    t = parse_xml(f'<w:t {nsdecls("w")}>目录将在打开文档后更新</w:t>')
    r2 = parse_xml(f'<w:r {nsdecls("w")}><w:t>目录将在打开文档后更新</w:t></w:r>')
    fldChar3 = parse_xml(f'<w:fldChar {nsdecls("w")} w:fldCharType="end"/>')
    run._element.append(fldChar1)
    run._element.append(instrText)
    run._element.append(fldChar2)
    run._element.append(t)
    run._element.append(fldChar3)

    doc.add_page_break()

    # ═══════════════════════════════════════════════════
    # RECIPIENT AND INTRODUCTION
    # ═══════════════════════════════════════════════════
    add_body_paragraph(doc, '致：上海景和商业保理有限公司', indent=False, bold=True)

    intro_text = (
        '【待补充：律所或法务部门名称】（以下简称"本出具主体"）受上海景和商业保理有限公司'
        '（以下简称"景和保理"或"委托方"）委托，就景和保理拟向浙江宏达供应链有限公司'
        '（以下简称"宏达供应链"）提供人民币5,000万元保理融资业务（以下简称"本事项"），'
        '依据截至2026年8月17日在中华人民共和国境内现行有效的法律、行政法规、部门规章、'
        '规范性文件及本出具主体已审阅的材料，出具本专项法律意见书。'
    )
    add_body_paragraph(doc, intro_text, indent=True)

    # ═══════════════════════════════════════════════════
    # SECTION 1: ENTRUSTED MATTER AND SCOPE
    # ═══════════════════════════════════════════════════
    add_heading(doc, '一、委托事项与意见范围', 1)

    add_body_paragraph(doc, '委托事项：就景和保理拟向宏达供应链提供5,000万元保理融资业务中涉及的应收账款真实性核验、'
        '债权转让效力、债务人抵销权、银行授信限制性条款合规性、保证与担保安排以及回款账户控制等七个法律问题，'
        '出具专项法律意见。', indent=True)

    add_body_paragraph(doc, '意见范围：本法律意见书涵盖本次保理融资交易中景和保理与宏达供应链之间的保理合同关系、'
        '宏达供应链与宁波海盛装备制造有限公司（以下简称"海盛装备"）之间的基础交易及应收账款转让关系、'
        '宏达供应链实际控制人周某的个人保证、杭州宏达物流有限公司（以下简称"宏达物流"）的关联担保，'
        '以及回款账户安排等法律事项。地域为中国大陆，期间为2026年1月至融资到期日。', indent=True)

    add_body_paragraph(doc, '排除事项：本法律意见书不对财务可行性、税务筹划、资产评估、审计意见、'
        '信用风险评级、宏观经营风险等非法律事项发表意见；不对中国大陆以外的法律适用问题发表意见。', indent=True)

    add_body_paragraph(doc, '法律基准日：2026年8月17日。', indent=True)

    # ═══════════════════════════════════════════════════
    # SECTION 2: REVIEWED MATERIALS
    # ═══════════════════════════════════════════════════
    add_heading(doc, '二、已审阅材料与核查方式', 1)

    add_body_paragraph(doc, '本法律意见书依据下列材料及核查结果形成。未提供或未能核验的材料已在相应分析中作为保留事项列明。', indent=True)

    # Materials table
    materials = [
        ['序号', '材料名称', '版本/日期', '核查状态'],
        ['1', '《年度采购框架协议》（宏达供应链与海盛装备）', '2026年1月5日', '委托方提供/未核验原件'],
        ['2', '送货单（部分）', '2026年1月至7月', '委托方提供/部分无公司印章'],
        ['3', '增值税专用发票', '2026年1月至7月', '委托方提供/待核验'],
        ['4', '对账单及邮件确认记录', '2026年1月至7月', '委托方提供/部分无公司盖章'],
        ['5', '宏达供应链银行授信合同', '日期未提供', '委托方陈述/未取得全文'],
        ['6', '宏达供应链负债及抵质押情况说明', '截至2026年7月底', '委托方陈述/未独立核验'],
        ['7', '周某资产情况说明', '截至2026年7月底', '委托方陈述/未独立核验'],
        ['8', '宏达物流公司章程', '日期未提供', '委托方提供/待核验'],
        ['9', '应收账款明细表（6,800万元）', '截至2026年7月31日', '委托方提供/待核验'],
    ]

    table = doc.add_table(rows=len(materials), cols=4)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row_data in enumerate(materials):
        for j, cell_text in enumerate(row_data):
            cell = table.rows[i].cells[j]
            cell.text = cell_text
            set_cell_font(cell, '楷体', Pt(11), bold=(i == 0))
            if i == 0:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Set column widths
    for row in table.rows:
        row.cells[0].width = Cm(1.5)
        row.cells[1].width = Cm(7)
        row.cells[2].width = Cm(4)
        row.cells[3].width = Cm(4)

    add_body_paragraph(doc, '核查方式：本出具主体对上述材料进行了书面审查，并通过公开法律数据库检索了相关法律、'
        '行政法规、司法解释及裁判案例。对标注"委托方陈述"或"未独立核验"的材料，本出具主体未进行独立核实，'
        '相关事实以委托方陈述为准，并在相应分析中予以保留。', indent=True)

    # ═══════════════════════════════════════════════════
    # SECTION 3: DECLARATIONS AND ASSUMPTIONS
    # ═══════════════════════════════════════════════════
    add_heading(doc, '三、声明、假设与保留', 1)

    add_body_paragraph(doc, '本出具主体假设所获材料真实、准确、完整，副本与原件一致，签署人具有相应权限；'
        '该等假设不替代对关键材料的必要核查。', indent=True)

    add_body_paragraph(doc, '对必须依赖专业机构结论的事项，本法律意见书仅作引述，不对其专业结论承担独立验证责任。', indent=True)

    add_body_paragraph(doc, '如关键事实、法律规则或监管口径在法律基准日后发生变化，本法律意见书的结论可能需要更新。', indent=True)

    add_body_paragraph(doc, '本事项的特别保留事项包括：（1）银行授信合同全文尚未取得，限制性条款的具体内容无法核验；'
        '（2）周某配偶的同意意向尚未明确；（3）宏达物流股东决定文件尚未形成；'
        '（4）海盛装备对应付账款的正式书面确认尚未取得；（5）部分送货单和对账单缺少公司印章或盖章。', indent=True)

    # ═══════════════════════════════════════════════════
    # SECTION 4: FACTS BACKGROUND
    # ═══════════════════════════════════════════════════
    add_heading(doc, '四、事实背景', 1)

    add_body_paragraph(doc, '（一）交易主体', indent=False, bold=True)

    add_body_paragraph(doc, '景和保理系一家依法设立并有效存续的商业保理公司，拟作为保理融资方（保理人）开展本次保理融资业务。'
        '宏达供应链系一家主要从事原材料采购和供应链服务的企业，为大型制造企业提供供应链服务，拟作为应收账款转让人（融资方）。'
        '海盛装备系一家装备制造企业，为本次保理融资业务中基础交易的债务人。'
        '宏达物流系宏达供应链的全资子公司，拟为本次融资提供连带责任保证。'
        '周某系宏达供应链的实际控制人，持有宏达供应链70%股权，拟为本次融资提供个人连带责任保证。', indent=True)

    add_body_paragraph(doc, '（二）基础交易与应收账款', indent=False, bold=True)

    add_body_paragraph(doc, '据委托方陈述，宏达供应链与海盛装备于2026年1月5日签订《年度采购框架协议》，约定宏达供应链向海盛装备'
        '供应钢材及工业零部件，年度预计交易金额约人民币1.2亿元。截至2026年7月31日，宏达供应链称其对海盛装备'
        '已形成应收账款共计人民币6,800万元，账期为90日至120日。宏达供应链提供了采购合同、送货单、增值税专用发票及对账单作为基础交易凭证。', indent=True)

    add_body_paragraph(doc, '（三）保理融资安排', indent=False, bold=True)

    add_body_paragraph(doc, '景和保理拟向宏达供应链提供人民币5,000万元保理融资，以宏达供应链对海盛装备享有的上述应收账款'
        '作为主要还款来源。景和保理要求海盛装备将应收账款直接支付至景和保理指定的监管账户。'
        '宏达供应链希望采用"隐蔽保理"方式，不立即通知海盛装备债权转让事宜。', indent=True)

    add_body_paragraph(doc, '（四）担保安排', indent=False, bold=True)

    add_body_paragraph(doc, '本次融资拟由周某提供个人连带责任保证。周某名下主要资产包括两套住宅及其持有的宏达供应链70%股权，'
        '其中一套住宅已抵押给银行，另一套住宅为夫妻共同所有，周某配偶尚未明确是否愿意签署相关文件。'
        '此外，宏达供应链计划由其全资子公司宏达物流提供连带责任保证，但宏达物流公司章程规定对外担保应当由股东决定，'
        '目前尚未形成正式股东决定文件。', indent=True)

    add_body_paragraph(doc, '（五）宏达供应链负债情况', indent=False, bold=True)

    add_body_paragraph(doc, '据委托方陈述，宏达供应链目前银行借款余额约人民币2.8亿元，已将部分存货、机器设备及银行账户'
        '质押或抵押给不同金融机构。其中一家银行的授信合同约定，宏达供应链未经银行同意不得新增重大融资或对外提供担保。'
        '景和保理尚未确认本次5,000万元融资是否会触发该授信合同中的限制性条款。', indent=True)

    doc.add_page_break()

    # ═══════════════════════════════════════════════════
    # SECTION 5: APPLICABLE LEGAL BASIS
    # ═══════════════════════════════════════════════════
    add_heading(doc, '五、适用法律依据', 1)

    add_body_paragraph(doc, '本法律意见书适用以下核心法律、行政法规及司法解释条文：', indent=True)

    # Law provisions with analysis
    provisions = [
        ('1. 《中华人民共和国民法典》第五百四十五条第二款规定："当事人约定非金钱债权不得转让的，不得对抗善意第三人。'
            '当事人约定金钱债权不得转让的，不得对抗第三人。"',
         '该款区分金钱债权与非金钱债权的禁止转让约定的效力。金钱债权的禁止转让约定不得对抗第三人，'
         '即基础合同中限制转让条款不影响金钱债权转让对保理人的效力。本事项中应收账款属于金钱债权。'),
        ('2. 《中华人民共和国民法典》第五百四十六条第一款规定："债权人转让债权，未通知债务人的，该转让对债务人不发生效力。"',
         '该款确立债权转让通知效力规则。未通知债务人时，债权转让对债务人不发生效力，债务人向原债权人清偿仍构成有效清偿。'
         '该规定直接影响隐蔽保理模式下的回款安全。'),
        ('3. 《中华人民共和国民法典》第五百四十九条规定："有下列情形之一的，债务人可以向受让人主张抵销：'
            '（一）债务人接到债权转让通知时，债务人对让与人享有债权，且债务人的债权先于转让的债权到期或者同时到期；'
            '（二）债务人的债权与转让的债权是基于同一合同产生。"',
         '该条赋予债务人在特定条件下对受让人主张抵销的权利。若债务人在收到转让通知时对原债权人享有到期债权，'
         '可向保理人主张抵销，从而减少保理人实际可收回的应收账款金额。'),
        ('4. 《中华人民共和国民法典》第七百六十三条规定："应收账款债权人与债务人虚构应收账款作为转让标的，'
            '与保理人订立保理合同的，应收账款债务人不得以应收账款不存在为由对抗保理人，但是保理人明知虚构的除外。"',
         '该条以债务人与债权人虚构应收账款为前提，确立债务人不得以应收账款不存在对抗保理人的规则，'
         '但以保理人不知情为条件。反向而言，保理人订立保理合同时对应收账款真实性负合理审查义务，'
         '若明知虚构仍订立合同，则不能依据该条对抗债务人。'),
        ('5. 《中华人民共和国民法典》第七百六十八条规定："应收账款债权人就同一应收账款订立多个保理合同，'
            '致使多个保理人主张权利的，已经登记的先于未登记的取得应收账款；均已经登记的，按照登记时间的先后顺序取得应收账款；'
            '均未登记的，由最先到达应收账款债务人的转让通知中载明的保理人取得应收账款；既未登记也未通知的，'
            '按照保理融资款或者服务报酬的比例取得应收账款。"',
         '该条确立保理登记优先规则。在隐蔽保理模式下，办理应收账款转让登记可以取得对抗后续保理人的优先权；'
         '未登记时，向债务人发出转让通知的时间先后亦影响权利顺位，是保护保理人融资安全的重要手段。'),
        ('6. 《中华人民共和国公司法》（2023年修订）第十五条规定："公司向其他企业投资或者为他人提供担保，'
            '按照公司章程的规定，由董事会或者股东会决议；公司章程对投资或者担保的总额及单项投资或者担保的数额'
            '有限额规定的，不得超过规定的限额。公司为公司股东或者实际控制人提供担保的，应当经股东会决议。'
            '前款规定的股东或者受前款规定的实际控制人支配的股东，不得参加前款规定事项的表决。'
            '该项表决由出席会议的其他股东所持表决权的过半数通过。"',
         '该条为公司对外担保的基本规则。公司为股东或实际控制人提供担保须经股东会决议。'
         '全资子公司为母公司提供担保属于为股东担保，须依法形成股东决定文件。债权人接受担保时应审查决议文件。'),
    ]

    for prov_text, application_text in provisions:
        add_law_provision(doc, f'{prov_text}{application_text}')

    doc.add_page_break()

    # ═══════════════════════════════════════════════════
    # SECTION 6: LEGAL ANALYSIS
    # ═══════════════════════════════════════════════════
    add_heading(doc, '六、法律分析', 1)

    # ── Issue 1 ──
    add_heading(doc, '（一）应收账款真实性核验风险', 2)

    add_body_paragraph(doc, '据委托方陈述，宏达供应链提供了采购合同、送货单、增值税专用发票及对账单作为应收账款的真实性证明。'
        '但存在以下问题：大部分送货单由海盛装备仓库人员签字，部分送货单未加盖公司印章；部分对账单仅经海盛装备采购部门'
        '员工邮件确认，无公司盖章；约1,200万元应收账款对应的货物仍处于质量复检阶段，海盛装备尚未明确确认最终应付金额。', indent=True)

    add_body_paragraph(doc, '根据《中华人民共和国民法典》第七百六十三条的规定，应收账款债权人与债务人虚构应收账款作为转让标的，'
        '与保理人订立保理合同的，应收账款债务人不得以应收账款不存在为由对抗保理人，但保理人明知虚构的除外。'
        '该条以保理人对虚构事实不知情为前提，反向要求保理人对应收账款真实性承担合理审查义务。'
        '商业保理监管要求保理企业受让应收账款应基于真实交易，对基础交易背景的真实性从严审查，'
        '对关联交易及账期异常等情形更应重点核查。', indent=True)

    add_body_paragraph(doc, '本事项中，送货单缺少公司印章但仅有仓库人员签字，存在签署权限不确定的风险。'
        '对账单仅通过邮件确认，缺乏正式书面盖章确认，证明力较弱。上述瑕疵虽不直接导致应收账款虚假，'
        '但降低了真实性核验的可靠性。特别是约1,200万元处于质量复检阶段的应收账款，因海盛装备尚未确认最终应付金额，'
        '该部分应收账款在金额确定性和可收回性方面存在显著风险。', indent=True)

    add_body_paragraph(doc, '综合现有材料，在取得海盛装备对应付账款的正式书面确认前，约1,200万元处于质量复检阶段的应收账款的真实性'
        '和确定性存在风险，不宜将该部分全额纳入融资基础资产。根据《最高人民法院关于适用〈中华人民共和国民法典〉'
        '合同编通则若干问题的解释》第四十九条的精神，受让人基于债务人对债权真实存在的确认受让债权后，'
        '债务人又以该债权不存在为由拒绝履行的，人民法院不予支持。因此，取得海盛装备对应收账款真实存在及金额的'
        '书面确认，可以有效锁定债务人的抗辩空间，对解决送货单、对账单签章瑕疵问题具有直接意义。'
        '其余约5,600万元应收账款的真实性基础较好，但建议补充取得加盖海盛装备公司印章的送货单和对账确认文件，'
        '以提高核验可靠性。', indent=True)

    # ── Issue 2 ──
    add_heading(doc, '（二）基础合同禁止转让条款对债权转让效力的影响', 2)

    add_body_paragraph(doc, '据委托方陈述，《年度采购框架协议》约定"未经买方书面同意，供应商不得向任何第三方转让本合同项下债权"，'
        '宏达供应链此前未就应收账款转让取得海盛装备书面同意。', indent=True)

    add_body_paragraph(doc, '根据《中华人民共和国民法典》第五百四十五条第二款的规定，当事人约定金钱债权不得转让的，'
        '不得对抗第三人。本事项中宏达供应链对海盛装备享有的应收账款属于金钱债权，基础合同中的禁止转让条款'
        '不影响该金钱债权转让对第三人（景和保理）的效力。即使未经海盛装备书面同意，宏达供应链将应收账款转让给景和保理'
        '在法律上仍然有效。', indent=True)

    add_body_paragraph(doc, '需注意的反方观点是：海盛装备可能以合同约定为由拒绝向景和保理付款或提出抗辩。'
        '虽然该抗辩在法律上不能成立（金钱债权禁止转让约定不得对抗第三人），但可能增加争议解决成本和时间。'
        '此外，根据《中华人民共和国民法典》第五百四十八条的规定，债务人对让与人的抗辩可以向受让人主张，'
        '海盛装备可能基于基础合同项下的其他事由（如货物质量问题）提出抗辩。', indent=True)

    add_body_paragraph(doc, '据此，基础合同中的禁止转让条款不影响金钱债权转让对景和保理的法律效力，'
        '宏达供应链将应收账款转让给景和保理合法有效。但为减少争议风险，建议在转让前或同时向海盛装备发出债权转让通知，'
        '或争取取得海盛装备对转让的书面同意。', indent=True)

    # ── Issue 3 ──
    add_heading(doc, '（三）债务人抵销权对保理融资回款的影响', 2)

    add_body_paragraph(doc, '据委托方陈述，海盛装备同时向宏达供应链销售生产设备，截至2026年7月底，宏达供应链尚欠海盛装备'
        '设备采购款约900万元。双方合同中未明确禁止抵销。海盛装备财务人员曾在邮件中表示"双方后续可以考虑把设备款'
        '和材料款统一核算"，但尚未签署正式抵销协议。', indent=True)

    add_body_paragraph(doc, '根据《中华人民共和国民法典》第五百四十九条的规定，债务人接到债权转让通知时，'
        '债务人对让与人享有债权且债务人的债权先于或同时于转让的债权到期的，债务人可以向受让人主张抵销。'
        '该规定的立法目的在于保护债务人在债权转让前已享有的抵销利益不因转让而丧失。', indent=True)

    add_body_paragraph(doc, '本事项中，海盛装备对宏达供应链享有约900万元设备款债权。如果在海盛装备收到债权转让通知时，'
        '该900万元设备款债权已到期或与转让的应收账款同时到期，海盛装备有权向景和保理主张抵销。'
        '这将导致景和保理实际可收回的应收账款金额减少约900万元，即从6,800万元减至约5,900万元，'
        '从而影响融资覆盖比例。', indent=True)

    add_body_paragraph(doc, '需注意，目前双方尚未签署正式抵销协议，海盛装备财务人员的邮件表述"可以考虑"尚不构成'
        '确定的抵销意思表示。但该邮件表明海盛装备具有行使抵销权的意向，景和保理应当将此风险纳入融资额度测算。', indent=True)

    add_body_paragraph(doc, '由此可以判断，海盛装备在满足法定条件下有权主张抵销约900万元设备款。'
        '建议在计算融资基础资产时扣减可能被抵销的金额，或要求宏达供应链就该部分提供额外担保。'
        '在隐蔽保理模式下，由于未通知海盛装备，其暂时无法向景和保理主张抵销，但一旦通知发出，该抵销权即可行使。', indent=True)

    # ── Issue 4 ──
    add_heading(doc, '（四）银行授信限制性条款的违约风险', 2)

    add_body_paragraph(doc, '据委托方陈述，宏达供应链目前银行借款余额约2.8亿元，其中一家银行的授信合同约定，'
        '宏达供应链未经银行同意不得新增重大融资或对外提供担保。景和保理尚未确认本次5,000万元保理融资'
        '是否会触发该限制性条款。', indent=True)

    add_body_paragraph(doc, '本问题的判断取决于银行授信合同中"重大融资"的具体定义、适用范围和触发门槛。'
        '由于银行授信合同全文尚未取得，本出具主体无法核验限制性条款的具体内容。'
        '根据合同法的基本原则，如果本次5,000万元保理融资被认定为授信合同项下的"新增重大融资"，'
        '且未经该银行同意，则可能构成对授信合同的违约，银行可能据此宣布贷款提前到期、'
        '要求提供额外担保或采取其他违约救济措施。', indent=True)

    add_body_paragraph(doc, '此外，宏达供应链已将部分存货、机器设备及银行账户质押或抵押给不同金融机构，'
        '如果银行宣布贷款提前到期并执行担保，可能进一步影响宏达供应链的偿债能力，间接影响本次保理融资的回款安全。', indent=True)

    add_body_paragraph(doc, '在现有材料范围内，由于银行授信合同全文尚未取得，限制性条款的具体内容无法核验，'
        '本出具主体就本事项暂无法判断。建议在融资放款前取得银行授信合同全文，审查限制性条款的具体内容、'
        '适用范围和触发条件，必要时与相关银行沟通确认本次融资是否构成违约。', indent=True)

    # ── Issue 5 ──
    add_heading(doc, '（五）实际控制人个人保证及夫妻共同财产问题', 2)

    add_body_paragraph(doc, '据委托方陈述，宏达供应链实际控制人周某愿意为本次融资提供个人连带责任保证。'
        '周某名下主要资产包括两套住宅及其持有的宏达供应链70%股权。其中一套住宅已抵押给银行，'
        '另一套住宅为夫妻共同所有，周某配偶尚未明确是否愿意签署相关文件。', indent=True)

    add_body_paragraph(doc, '就周某个人连带责任保证的效力而言，周某作为完全民事行为能力人，'
        '以其个人名义提供连带责任保证的意思表示有效，保证合同自双方签署后成立并生效。'
        '周某持有的宏达供应链70%股权以及其名下未设定抵押的财产可以作为保证责任的一般责任财产。', indent=True)

    add_body_paragraph(doc, '就夫妻共同所有住宅的抵押问题而言，根据《中华人民共和国民法典》关于夫妻共同财产的规定，'
        '夫妻在婚姻关系存续期间所得的财产属于夫妻共同财产。以夫妻共同财产设定抵押，应当取得配偶的同意。'
        '如果周某配偶不同意签署相关文件，则该套住宅上的抵押权可能无法有效设立。'
        '在司法实践中，夫妻一方未经另一方同意以共同财产设定抵押，另一方可以主张抵押无效或撤销。', indent=True)

    add_body_paragraph(doc, '此外，周某已抵押给银行的第一套住宅，其剩余价值（如有）受制于在先抵押权人的优先受偿权。'
        '周某持有的70%股权可能已质押给其他金融机构，需核查该股权是否存在权利负担。', indent=True)

    add_body_paragraph(doc, '结合上述情况，周某的个人连带责任保证本身有效，但以夫妻共同所有住宅设定抵押的效力'
        '取决于配偶是否同意。在周某配偶明确签署同意函前，该套住宅不能作为可靠的担保物。'
        '建议督促周某配偶尽快签署同意函；如配偶不同意，应重新评估周某保证的实际担保价值，'
        '并考虑要求增加其他担保措施。', indent=True)

    # ── Issue 6 ──
    add_heading(doc, '（六）关联公司担保的决议程序合规性', 2)

    add_body_paragraph(doc, '据委托方陈述，宏达供应链计划由其全资子公司宏达物流提供连带责任保证。'
        '宏达物流公司章程规定对外担保应当由股东决定，但目前尚未形成正式股东决定文件。', indent=True)

    add_body_paragraph(doc, '根据《中华人民共和国公司法》（2023年修订）第十五条的规定，公司为公司股东或者实际控制人'
        '提供担保的，应当经股东会决议。宏达物流为宏达供应链的子公司，宏达供应链是宏达物流的唯一股东。'
        '宏达物流为宏达供应链提供担保，属于"为公司股东提供担保"的情形，依法应当经股东会决议。'
        '由于宏达物流是全资子公司，宏达供应链作为唯一股东可以作出股东决定，但必须形成正式的书面股东决定文件。', indent=True)

    add_body_paragraph(doc, '在司法实践中，最高人民法院在（2021）最高法民再232号公报案例中明确认定，'
        '公司为股东提供担保未经股东会决议，债权人未审查股东会决议的，不构成善意相对人，担保合同无效。'
        '该案中，最高人民法院将担保合同无效后的公司赔偿责任调整为债务人不能清偿部分的二分之一。'
        '多个下级法院案例亦遵循同一裁判路径，认定未经决议的公司对外担保对担保公司不发生效力。', indent=True)

    add_body_paragraph(doc, '需注意的反方观点是：司法实践中存在公司为其全资子公司开展经营活动提供担保时，'
        '未经决议亦可能不当然免除担保责任的例外情形。但本案担保方向相反——系全资子公司宏达物流为母公司宏达供应链'
        '（即其股东）提供担保，依法属于公司为股东提供担保，应当经股东会决议，不适用上述例外情形。'
        '此外，根据《中华人民共和国民法典》第八十五条和《最高人民法院关于适用〈中华人民共和国公司法〉'
        '若干问题的规定（四）》第六条的规定，公司依据决议与善意相对人形成的民事法律关系不受影响；'
        '但景和保理作为专业保理机构，应当知道公司为股东担保须经股东会决议，'
        '若未审查股东决定文件，难以被认定为善意相对人。', indent=True)

    add_body_paragraph(doc, '据此，在宏达物流未形成正式股东决定文件前，该担保存在被认定对宏达物流不发生效力的风险。'
        '建议在放款前要求宏达物流出具经合法程序形成的股东决定文件，并由景和保理进行形式审查。'
        '如股东决定文件无法在放款前取得，应将该担保视为待生效条件，并相应调整融资方案或要求补充其他担保措施。', indent=True)

    # ── Issue 7 ──
    add_heading(doc, '（七）隐蔽保理模式下的回款账户控制风险', 2)

    add_body_paragraph(doc, '据委托方陈述，景和保理要求海盛装备将应收账款直接支付至景和保理指定的监管账户。'
        '但宏达供应链希望采用"隐蔽保理"方式，不立即通知海盛装备债权转让事宜，担心通知可能影响双方商业合作。'
        '景和保理则担心不通知可能导致海盛装备继续向宏达供应链付款，从而影响融资安全。', indent=True)

    add_body_paragraph(doc, '根据《中华人民共和国民法典》第五百四十六条第一款的规定，债权人转让债权，'
        '未通知债务人的，该转让对债务人不发生效力。在隐蔽保理模式下，景和保理不通知海盛装备债权转让事宜，'
        '债权转让对海盛装备不发生效力。海盛装备向宏达供应链付款仍构成有效清偿，景和保理无法直接向海盛装备主张付款。'
        '这就产生了回款不受控制的风险：海盛装备可能继续向宏达供应链付款，而宏达供应链可能挪用该款项。'
        '但根据《最高人民法院关于适用〈中华人民共和国民法典〉合同编通则若干问题的解释》第四十八条的规定，'
        '受让人直接起诉债务人请求履行债务，人民法院经审理确认债权转让事实的，债权转让自起诉状副本送达时'
        '对债务人发生效力，为保理人在隐蔽保理模式下保留了一条通过诉讼通知债务人并锁定债务履行的路径。', indent=True)

    add_body_paragraph(doc, '根据《中华人民共和国民法典》第七百六十八条规则，已登记的保理人先于未登记的取得应收账款。'
        '因此，即使采用隐蔽保理模式，景和保理仍应在中国人民银行征信中心动产融资统一登记公示系统办理应收账款转让登记，'
        '以取得对抗后续保理人的优先权。', indent=True)

    add_body_paragraph(doc, '为平衡隐蔽保理的商业需求与融资安全，建议采取以下综合措施：'
        '第一，办理应收账款转让登记，取得优先权保护；'
        '第二，与宏达供应链签订账户监管协议，约定海盛装备回款必须进入指定监管账户，'
        '由景和保理对账户资金进行监控和划转控制；'
        '第三，约定触发式通知机制，在宏达供应链违约、回款异常或出现其他约定情形时，'
        '景和保理有权立即向海盛装备发出债权转让通知，使债权转让对海盛装备发生效力；'
        '第四，要求宏达供应链承诺不以任何方式免除或减免海盛装备的债务，不擅自延长付款期限。', indent=True)

    add_body_paragraph(doc, '综合来看，隐蔽保理模式下未通知债务人时债权转让对债务人不发生效力，存在债务人继续向原债权人付款的风险。'
        '通过应收账款转让登记、账户监管协议和触发式通知机制可以有效降低该风险。'
        '在上述措施到位的前提下，隐蔽保理模式可以在一定程度上保障景和保理的融资安全。', indent=True)

    doc.add_page_break()

    # ═══════════════════════════════════════════════════
    # SECTION 7: RISK MATRIX
    # ═══════════════════════════════════════════════════
    add_heading(doc, '七、风险分级与应对建议', 1)

    risks = [
        ['风险事项', '等级', '触发条件', '建议措施', '责任人/期限'],
        ['1,200万元应收账款\n真实性不确定', '高', '货物质量复检结果\n未确定，应付金额\n未经海盛装备确认', '取得海盛装备书面\n确认函；暂不将该\n部分纳入融资基础', '景和保理\n放款前'],
        ['债务人抵销权\n（900万元）', '中', '海盛装备在收到\n转让通知后主张\n抵销设备款', '在融资额度中扣减\n可能被抵销金额；\n要求额外担保', '景和保理\n额度测算阶段'],
        ['关联担保决议\n缺失', '高', '宏达物流未形成\n股东决定文件', '要求宏达物流出具\n股东决定文件并\n进行形式审查', '宏达供应链\n放款前'],
        ['保证人配偶\n未同意', '中', '周某配偶拒绝签署\n共同财产抵押\n同意函', '督促配偶签署；\n如不同意则重新\n评估保证价值', '周某\n放款前'],
        ['隐蔽保理\n回款风险', '中', '海盛装备继续向\n宏达供应链付款', '办理转让登记；\n签订账户监管协议；\n约定触发式通知', '景和保理\n放款前'],
        ['银行授信\n限制性条款', '待核', '本次融资被认定\n为"新增重大融资"', '取得授信合同全文\n并审查限制条款；\n必要时与银行沟通', '景和保理\n放款前'],
        ['禁止转让条款\n争议风险', '低', '海盛装备以合同\n约定提出抗辩', '发出转让通知或\n取得书面同意', '景和保理\n放款前'],
    ]

    table2 = doc.add_table(rows=len(risks), cols=5)
    table2.style = 'Table Grid'
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row_data in enumerate(risks):
        for j, cell_text in enumerate(row_data):
            cell = table2.rows[i].cells[j]
            cell.text = cell_text
            set_cell_font(cell, '楷体', Pt(10), bold=(i == 0))
            if i == 0:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)

    # Set column widths
    for row in table2.rows:
        row.cells[0].width = Cm(3.5)
        row.cells[1].width = Cm(1.5)
        row.cells[2].width = Cm(4)
        row.cells[3].width = Cm(4)
        row.cells[4].width = Cm(2.5)

    # Mark table header row to repeat
    header_row = table2.rows[0]._tr
    trPr = header_row.get_or_add_trPr()
    tblHeader = parse_xml(f'<w:tblHeader {nsdecls("w")}/>')
    trPr.append(tblHeader)

    doc.add_page_break()

    # ═══════════════════════════════════════════════════
    # SECTION 8: CONCLUSIONS
    # ═══════════════════════════════════════════════════
    add_heading(doc, '八、结论意见', 1)

    add_body_paragraph(doc, '基于本法律意见书列明的材料、事实、核查结果和保留事项，本出具主体认为：', indent=True)

    conclusions = [
        '1. 本次保理融资交易整体可以推进，但须在放款前完成以下关键事项的核验与补充：'
        '取得海盛装备对应付账款的正式书面确认（特别是1,200万元质量复检部分）、取得银行授信合同全文并审查限制性条款、'
        '督促周某配偶签署夫妻共同财产抵押同意函、要求宏达物流出具经合法程序形成的股东决定文件。',
        '2. 基础合同中的禁止转让条款不影响金钱债权转让对景和保理的法律效力，宏达供应链将应收账款转让给景和保理合法有效。'
        '建议在转让后向海盛装备发出债权转让通知或争取取得其书面同意，以减少争议风险。',
        '3. 海盛装备在满足《中华人民共和国民法典》第五百四十九条规定条件下有权主张抵销约900万元设备款，'
        '建议在计算融资基础资产时扣减可能被抵销的金额或要求宏达供应链就该部分提供额外担保。',
        '4. 周某的个人连带责任保证本身有效，但以夫妻共同所有住宅设定抵押的效力取决于配偶是否同意；'
        '宏达物流的关联担保在未形成正式股东决定文件前存在被认定不发生效力的风险。'
        '上述担保措施的完善是放款前必须完成的条件。',
        '5. 隐蔽保理模式下，建议采取"应收账款转让登记+账户监管协议+触发式通知机制"的综合方案，'
        '在满足宏达供应链商业需求的同时保障景和保理的融资安全。',
    ]
    for c in conclusions:
        add_body_paragraph(doc, c, indent=True)

    # ═══════════════════════════════════════════════════
    # SECTION 9: USAGE AND VALIDITY
    # ═══════════════════════════════════════════════════
    add_heading(doc, '九、用途与有效期', 1)

    add_body_paragraph(doc, '本法律意见书以已取得材料真实、准确、完整，副本与原件一致，签署人具有相应权限为基础；'
        '该等假设不替代对关键材料的必要核查。', indent=True)

    add_body_paragraph(doc, '本事项的保留事项包括：银行授信合同全文尚未取得，限制性条款的具体内容无法核验；'
        '周某配偶的同意意向尚未明确；宏达物流股东决定文件尚未形成；海盛装备对应付账款的正式书面确认尚未取得。'
        '上述事项的核实结果将直接影响相应结论的成立。', indent=True)

    add_body_paragraph(doc, '本法律意见书仅供上海景和商业保理有限公司为本次保理融资业务的内部审批和交易决策使用。'
        '本法律意见书基于2026年8月17日的事实与法律状态形成；相关事实、交易结构、监管规则或审批状态发生变化时，'
        '应对受影响结论重新评估。', indent=True)

    add_body_paragraph(doc, '本法律意见书为待复核初稿，未经有权律师或专业人员实质复核签发，不得作为正式法律意见对外使用。', indent=True)

    doc.add_page_break()

    # ═══════════════════════════════════════════════════
    # SIGNATURE PAGE
    # ═══════════════════════════════════════════════════
    add_body_paragraph(doc, '（本页无正文，为本法律意见书签署页。）', indent=False)

    for _ in range(6):
        add_empty_paragraph(doc)

    add_body_paragraph(doc, '出具主体：【待补充：律所或法务部门名称】', indent=False)
    add_empty_paragraph(doc)
    add_body_paragraph(doc, '经办律师：【待补充：律师姓名】', indent=False)
    add_empty_paragraph(doc)
    add_body_paragraph(doc, '执业证号：【待补充：执业证号】', indent=False)
    add_empty_paragraph(doc)
    add_body_paragraph(doc, '日期：2026年8月17日', indent=False)

    # ── Save ──
    output_path = Path(output_path).resolve()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    print(f'DOCX saved to: {output_path}')

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description='生成通用专项法律意见书 DOCX 样稿')
    parser.add_argument('--output', default=DEFAULT_OUTPUT_NAME, help='输出 DOCX 路径')
    args = parser.parse_args()
    generate(args.output)
